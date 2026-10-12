#!/usr/bin/env python3
"""Analyze the 2024-09-24 audit batch without generating or changing data.

Outputs a detailed XLSX and JSON plan for human approval. The analyzer keeps
source+target decisions conservative and records every shared-scene spillover.
"""
from __future__ import annotations

import argparse
import difflib
import json
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

SOURCE_TOKENS = ("source", "原图", "原source", "重新设计source", "重新生成source", "空间不足", "画面拥挤", "物体太多太杂", "两行一样", "同一模")
COUNT_SUBTYPES = {"1_to_2", "2_to_1", "2_to_3", "3_to_2", "3_to_1", "1_to_3"}
RELATION_SUBTYPES = {"occlusion_state", "containment"}
MOVE_SUBTYPES = {"absolute_horizontal", "absolute_vertical"}
ATTRIBUTE_SUBTYPES = {"material", "pattern", "functional_state"}


def load_manifest(path: Path) -> dict[str, dict[str, Any]]:
    return {r["pair_id"]: r for r in (json.loads(l) for l in path.open(encoding="utf-8") if l.strip())}


def load_audit(path: Path) -> list[dict[str, Any]]:
    wb = __import__("openpyxl").load_workbook(path, read_only=True, data_only=True)
    ws = wb["新增问题汇总"]
    headers = [c.value for c in ws[1]]
    return [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]


def entity_info(pair: dict[str, Any]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    iv = pair.get("intervention", {})
    eid = str(iv.get("entity_id", ""))
    entities = pair.get("source", {}).get("cot", {}).get("entities", [])
    target = next((e for e in entities if str(e.get("id")) == eid), {})
    return target, entities


def count_ambiguity(pair: dict[str, Any], issue: str) -> tuple[str, str]:
    iv = pair.get("intervention", {})
    entity, entities = entity_info(pair)
    count = entity.get("count")
    text = issue.lower()
    if str(iv.get("field")) == "count" or pair.get("subtype") in COUNT_SUBTYPES:
        if isinstance(count, int) and count > 1:
            return (
                "高",
                f"entity_id={iv.get('entity_id')} 是 count={count} 的组级实体；必须明确按 individual instances 计数，不按沓/摞/行计数。"
            )
        if any(x in text for x in ("一沓", "两沓", "三沓", "一行", "两行", "四个", "数量不对")):
            return ("中", "审计文字显示组/堆/行与individual count可能混淆；source需显式规定唯一目标组和最终个体总数。")
    if str(iv.get("entity_id", "")).endswith("_1") and isinstance(count, int) and count > 1:
        return ("中", "entity_id以_1结尾但仍是组级count实体；不能自动解释为第一个独立实例。")
    return ("低", "未发现明显的组级计数歧义；仍需按COT entity_id和count人工核对。")


def classify(pair: dict[str, Any], issue: str, explicit_source: bool) -> tuple[str, str, str]:
    subtype = str(pair.get("subtype", ""))
    low = issue.lower()
    source_reason = []
    if explicit_source or any(t.lower() in low for t in SOURCE_TOKENS):
        source_reason.append("用户指出source初始结构、数量、遮挡、容器或可用空间问题")
    if subtype in COUNT_SUBTYPES and any(t in low for t in ("source", "数量不对", "两行", "两沓", "四个", "没有三个", "只保留一层")):
        source_reason.append("数量任务的source old状态/唯一目标区域不可信")
    if subtype in RELATION_SUBTYPES and any(t in low for t in ("source", "遮盖关系", "容器关系", "看不见", "里面", "外面", "移入", "移出")):
        source_reason.append("关系任务依赖source初始空间关系和可见性")
    if subtype in MOVE_SUBTYPES and any(t in low for t in ("空间", "位置", "妨碍", "没有", "移动", "桌篼", "留出")):
        source_reason.append("移动任务需要目标方向有净空，且无关物体不被挤动")
    if subtype in ATTRIBUTE_SUBTYPES and any(t in low for t in ("source", "状态不对", "物体似乎不适合")):
        source_reason.append("属性任务的source对象/初始状态不可信")
    if source_reason:
        return "source+target", "；".join(source_reason), "高"
    if subtype in RELATION_SUBTYPES:
        return "target-only", "source初始条件暂可用，优先用明确的已有对象/容器关系prompt修复target；若复核发现source不成立再升级source+target。", "中"
    if subtype in COUNT_SUBTYPES:
        return "target-only", "source暂未被明确否定，使用同一目标组/同一摞内部的individual count prompt修复target。", "中"
    if subtype in MOVE_SUBTYPES:
        return "target-only", "主要是target保留方向、旧位置清除或不动无关物体的问题；若确认净空不足再升级source+target。", "中"
    return "target-only", "主要是target原子编辑执行失败；保持source，强化target prompt和must_preserve约束。", "中"


def measures(pair: dict[str, Any], issue: str, scope: str) -> tuple[str, str, str, str]:
    iv = pair.get("intervention", {})
    subtype = str(pair.get("subtype", "")); field = str(iv.get("field", "")); old = iv.get("old_value"); new = iv.get("new_value")
    entity, _ = entity_info(pair); eid = iv.get("entity_id"); category = str(entity.get("category", "object")).replace("_", " "); count = entity.get("count")
    preservation = "仅允许 intervention.entity_id 的 intervention.field 发生一个原子变化；所有 must_preserve 实体、数量、位置、尺寸、方向、颜色、材质、背景和相机保持。"
    if subtype in COUNT_SUBTYPES or field == "count":
        if isinstance(old, int) and isinstance(new, int) and new > old:
            target = f"只在 entity_id={eid} 对应的同一目标区域/同一摞内部增加恰好 {new-old} 个 individual {category}；不生成第二摞、第二行、第二层或新区域；原有{old}个保持原位和细节；最终恰好{new}个。"
        else:
            target = f"只在 entity_id={eid} 对应的同一目标区域/同一摞内部移除恰好 {int(old)-int(new) if isinstance(old,int) and isinstance(new,int) else 1} 个 individual {category}；幸存者原位原细节；最终恰好{new}个；不改变无关物体。"
        source = "" if scope == "target-only" else "Source只保留一个明确目标区域/一摞，old_value个individual实例清晰可数；删除重复行、重复区域和无关同类物体；为增加方向预留空位。"
    elif subtype in RELATION_SUBTYPES:
        target = f"只改变 entity_id={eid} 的{field}关系，从{old}到{new}；使用已有对象，不新增、不复制、不删除对象；容器/遮挡物保持原位、尺寸、形状和属性；关系必须清晰可见。"
        source = "" if scope == "target-only" else "Source中只保留一个清晰的容器/参照物和目标物，初始关系与source COT完全一致；目标放入后仍可见，目标移出后仍可见；留出无遮挡空间。"
    elif subtype in MOVE_SUBTYPES:
        target = f"只把 entity_id={eid} 从{old}移动到{new}；保持数量、身份、尺寸、方向和细节；清除旧位置残留；所有无关对象原地不动，不旋转、不缩放、不切割。"
        source = "" if scope == "target-only" else "Source在目标方向预留完整净空，目标不会碰撞、穿透、遮挡或挤动无关物体；保留原方向参照。"
    elif subtype in ATTRIBUTE_SUBTYPES:
        target = f"只改变 entity_id={eid} 的{field}从{old}到{new}；变化必须在目标物体上肉眼明显；位置、数量、轮廓、身份及其他属性完全不变。"
        source = "" if scope == "target-only" else "Source目标物体轮廓、初始状态和光照清晰，且没有同类干扰物。"
    else:
        target = f"只执行 intervention.field={field} 的单原子变化{old}->{new}，其他一切保持。"
        source = "" if scope == "target-only" else "简化source布局并确保COT实体可见、可辨。"
    if any(x in issue for x in ("不明显", "不明确", "看不出")): target += "差异必须在并排比较中一眼可辨，模糊变化视为失败。"
    if any(x in issue for x in ("蜡", "流蜡")): target += "保留幸存蜡烛的滴蜡、流蜡、污渍和燃烧细节，不得换成崭新物体。"
    if any(x in issue for x in ("两行", "两沓", "四个", "太多", "满满")): source += " 只生成任务所需的一组对象，不生成重复行、重复堆或额外同类对象。"
    return source, target, preservation, "人工验收：source old状态/数量；target new状态/数量；无新增/消失；无关对象原位；mask局部；不合格不融合。"


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__); ap.add_argument('--audit-xlsx',required=True,type=Path); ap.add_argument('--manifest',required=True,type=Path); ap.add_argument('--output-xlsx',required=True,type=Path); ap.add_argument('--output-jsonl',required=True,type=Path); args=ap.parse_args()
    rows=load_audit(args.audit_xlsx); manifest=load_manifest(args.manifest); scene_members=defaultdict(list)
    for pid,r in manifest.items(): scene_members[r['source_scene_id']].append(pid)
    dedup={}
    for row in rows:
        pid=row.get('pair_id_manifest') or row.get('pair_id_raw')
        if pid not in manifest: continue
        if pid not in dedup or row.get('用户明确source重生成')=='是': dedup[pid]=row
    plans=[]
    for pid,row in sorted(dedup.items()):
        pair=manifest[pid]; issue=str(row.get('本次审计问题') or ''); explicit=row.get('用户明确source重生成')=='是'
        scope,reason,conf=classify(pair,issue,explicit); amb,amb_reason=count_ambiguity(pair,issue); src,tgt,pres,val=measures(pair,issue,scope)
        plans.append({'pair_id':pid,'source_scene_id':pair['source_scene_id'],'scene_pair_ids':scene_members[pair['source_scene_id']],'scene_pair_count':len(scene_members[pair['source_scene_id']]),'subtype':pair.get('subtype'),'intervention':pair.get('intervention'),'source_cot_text':pair['source'].get('cot_text'),'target_cot_text':pair['target'].get('cot_text'),'original_target_prompt':pair['target'].get('prompt'),'user_issue':issue,'user_source_marked':explicit,'repair_scope':scope,'scope_reason':reason,'confidence':conf,'entity_id':pair.get('intervention',{}).get('entity_id'),'entity_group_count':entity_info(pair)[0].get('count'),'entity_count_ambiguity':amb,'entity_ambiguity_reason':amb_reason,'source_measure':src,'target_measure':tgt,'preservation_measure':pres,'validation_measure':val,'approval_status':'待人工审核','generation_status':'未生成','merge_status':'未融合'})
    args.output_jsonl.parent.mkdir(parents=True,exist_ok=True); args.output_jsonl.write_text('\n'.join(json.dumps(x,ensure_ascii=False,sort_keys=True) for x in plans)+'\n',encoding='utf-8')
    wb=Workbook(); ws=wb.active; ws.title='修补分析'; headers=['序号','pair_id','source_scene_id','scene_pair_count','scene_pair_ids','subtype','intervention_field','old_value','new_value','entity_id','entity_group_count','entity_count_ambiguity','entity_ambiguity_reason','用户问题','用户明确source','repair_scope','scope_reason','confidence','Source修改措施','Target修改措施','单原子/must_preserve','验收检查','原始TargetPrompt','审批状态','生成状态','融合状态']; ws.append(headers)
    for i,p in enumerate(plans,1):
        iv=p['intervention'] or {}; ws.append([i,p['pair_id'],p['source_scene_id'],p['scene_pair_count'],';'.join(p['scene_pair_ids']),p['subtype'],iv.get('field'),iv.get('old_value'),iv.get('new_value'),p['entity_id'],p['entity_group_count'],p['entity_count_ambiguity'],p['entity_ambiguity_reason'],p['user_issue'],'是' if p['user_source_marked'] else '否',p['repair_scope'],p['scope_reason'],p['confidence'],p['source_measure'],p['target_measure'],p['preservation_measure'],p['validation_measure'],p['original_target_prompt'],p['approval_status'],p['generation_status'],p['merge_status']])
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F4E78'); c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    for rr in ws.iter_rows(min_row=2):
        for c in rr: c.alignment=Alignment(vertical='top',wrap_text=True)
    widths=[8,60,42,15,100,22,30,15,15,22,16,18,70,75,15,18,80,14,90,110,80,80,100,16,16,16]
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    tab=Table(displayName='RepairAnalysis20240924',ref=f'A1:Z{ws.max_row}'); tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True,showColumnStripes=False); ws.add_table(tab)
    st=wb.create_sheet('统计'); st.append(['统计项','数量']); st.append(['唯一pair',len(plans)]); st.append(['source+target建议',sum(p['repair_scope']=='source+target' for p in plans)]); st.append(['target-only建议',sum(p['repair_scope']=='target-only' for p in plans)]); st.append(['涉及共享source pair',sum(p['scene_pair_count']>1 for p in plans)]); st.append([]); st.append(['entity歧义','数量'])
    for k,v in Counter(p['entity_count_ambiguity'] for p in plans).items(): st.append([k,v])
    args.output_xlsx.parent.mkdir(parents=True,exist_ok=True); wb.save(args.output_xlsx); print(json.dumps({'output_xlsx':str(args.output_xlsx),'output_jsonl':str(args.output_jsonl),'unique_pairs':len(plans),'source_target':sum(p['repair_scope']=='source+target' for p in plans),'target_only':sum(p['repair_scope']=='target-only' for p in plans)},ensure_ascii=False,indent=2))
if __name__=='__main__': main()
