#!/usr/bin/env python3
"""Analyze source-regeneration candidates without changing data or calling APIs.

Reads the first-step XLSX and final manifest, then adds a conservative
recommendation column based on shared-scene risk, source count mismatches,
containment/occlusion feasibility, and spatial room. It never generates images,
modifies the manifest, or deletes pairs.
"""
from __future__ import annotations

import argparse
import json
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

SOURCE_WORDS = (
    "source", "Source", "source图", "Source图", "重新生成source", "重新生成 source",
    "重新设计source", "重新设计 source", "修改source", "修改 source",
)
HARD_SOURCE_REASONS = (
    "source有", "Source有", "source没有", "Source没有", "source并非", "Source并非",
    "source不是", "Source不是", "source图", "Source图", "source太复杂", "Source太复杂",
    "重新生成source", "重新生成 source", "重新设计source", "重新设计 source",
    "修改source", "修改 source", "source容器", "Source容器", "source前后", "Source前后",
)


def load_manifest(path: Path) -> dict[str, dict[str, Any]]:
    result = {}
    for line in path.open(encoding="utf-8"):
        if line.strip():
            row = json.loads(line)
            result[row["pair_id"]] = row
    return result


def read_summary(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["问题汇总"]
    headers = [cell.value for cell in ws[1]]
    return headers, [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]


def entities(record: dict[str, Any]) -> list[dict[str, Any]]:
    return record.get("source", {}).get("cot", {}).get("entities", [])


def source_recommendation(record: dict[str, Any], issue: str, scene_size: int) -> tuple[str, str, str]:
    subtype = str(record.get("subtype", ""))
    intervention = record.get("intervention", {})
    field = str(intervention.get("field", ""))
    issue_lower = issue.lower()
    reasons: list[str] = []
    confidence = "低"

    explicit = any(token.lower() in issue_lower for token in HARD_SOURCE_REASONS)
    if explicit:
        reasons.append("用户已明确指出source结构/数量/复杂度/可见性问题")
        confidence = "高"

    if subtype in {"occlusion_state", "containment"}:
        reasons.append("该任务依赖source中的初始空间关系；若初始关系不清晰，target单独重生成无法可靠修复")
        confidence = "高" if explicit else "中"
    if subtype in {"1_to_2", "2_to_1", "2_to_3", "3_to_2", "3_to_1", "1_to_3"} or field == "count":
        reasons.append("数量任务需要source中目标个体数量、层级和可数边界清晰；复杂/重复/干扰实体会导致错数")
        confidence = "高" if explicit else "中"
    if subtype in {"absolute_horizontal", "absolute_vertical"}:
        reasons.append("空间移动任务需要source在目标方向预留足够空白，且不能与其他物体重叠")
        confidence = "高" if explicit else "中"
    if scene_size > 1:
        reasons.append(f"该source_scene_id共享{scene_size}个pair；若重生成source，必须同步评估同scene全部pair")
    if "material" in subtype or "functional_state" in subtype or "pattern" in subtype:
        reasons.append("source对象需轮廓、材质/状态和光照清晰；否则target改动难以判定")
        if not confidence == "高":
            confidence = "中"

    if not reasons:
        reasons.append("当前证据更像target生成错误；可先只重生成target，暂不改source")
        recommendation = "target-only优先"
    else:
        recommendation = "建议source+target"
    return recommendation, confidence, "；".join(reasons)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--summary-xlsx", required=True, type=Path)
    ap.add_argument("--manifest", required=True, type=Path)
    ap.add_argument("--output-xlsx", required=True, type=Path)
    args = ap.parse_args()

    headers, rows = read_summary(args.summary_xlsx)
    manifest = load_manifest(args.manifest)
    scene_counts = Counter(str(r.get("source_scene_id", "")) for r in manifest.values())

    extra_source = 0
    for row in rows:
        pid = row.get("pair_id_manifest") or row.get("pair_id_raw")
        record = manifest.get(pid, {})
        issue = str(row.get("用户审计问题", ""))
        marked = str(row.get("用户明确标记source重生成", "")) == "是"
        if marked:
            recommendation, confidence, reason = "用户已标记source+target", "用户标记", "保留用户source决定；target需在新source上重新生成"
        else:
            recommendation, confidence, reason = source_recommendation(record, issue, scene_counts.get(str(record.get("source_scene_id", "")), 0))
            if recommendation == "建议source+target":
                extra_source += 1
        row["建议是否重生成source"] = recommendation
        row["建议置信度"] = confidence
        row["建议source原因"] = reason
        row["共享source_scene_id"] = record.get("source_scene_id", "")
        row["同scene pair数"] = scene_counts.get(str(record.get("source_scene_id", "")), 0)

    output_headers = headers + ["建议是否重生成source", "建议置信度", "建议source原因", "共享source_scene_id", "同scene pair数"]
    wb = load_workbook(args.summary_xlsx)
    ws = wb["问题汇总"]
    ws.delete_rows(1, ws.max_row)
    ws.append(output_headers)
    for row in rows:
        ws.append([row.get(header, "") for header in output_headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = [8,58,58,28,24,30,16,16,70,18,28,50,20,24,14,90,45,14]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    if ws.max_row > 1:
        table = Table(displayName="AuditIssueSummaryWithSourceAnalysis", ref=f"A1:R{ws.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        ws.add_table(table)

    stats = wb["统计"]
    stats.append([])
    stats.append(["本步新增分析", "数量"])
    stats.append(["用户标记source+target", sum(str(r.get("用户明确标记source重生成", "")) == "是" for r in rows)])
    stats.append(["额外建议source+target", extra_source])
    stats.append(["建议target-only优先", sum(r.get("建议是否重生成source") == "target-only优先" for r in rows)])
    stats.append(["共享source_scene_id涉及问题pair数", sum(int(r.get("同scene pair数") or 0) > 1 for r in rows)])

    args.output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output_xlsx)
    print(json.dumps({"output": str(args.output_xlsx), "rows": len(rows), "extra_source_recommendations": extra_source}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
