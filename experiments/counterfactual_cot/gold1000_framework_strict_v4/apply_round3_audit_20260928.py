#!/usr/bin/env python3
"""Apply the round3 human audit (2026-09-28) and build the replacement set for the 1000 pairs.

Rules
  - round2 audit deleted pairs stay deleted
  - the 6 pairs listed by the user on round3 are deleted (no further repair)
  - every other round3 regenerated pair passes
  - every round2 accepted pair passes

Outputs
  - repair_status_20260928_final.jsonl   final decision for all 49 audited pairs
  - repair_status_20260928_final.xlsx    statistics + lists
  - final_replacement_20260928.jsonl     42 records ready to replace into the 1000
  - final_replacement_20260928_summary.json
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path("/group/40016/niightli/BAGEL/experiments/counterfactual_cot/gold1000_framework_strict_v4")
ROUND2_DECISIONS = ROOT / "round3_audit_decisions_20260928.jsonl"
ROUND2_SANDBOX = ROOT / "repair_batch_20260924_round2_strict_atomic_v4"
ROUND2_MANIFEST = ROUND2_SANDBOX / "manifest.jsonl"
ROUND3_SANDBOX = ROOT / "repair_batch_20260928_round3"
ROUND3_MANIFEST = ROUND3_SANDBOX / "manifest.jsonl"
ORIGINAL_MANIFEST = ROOT / "full_1000_20260918" / "gold1000.generated.jsonl"

# round3 human audit: delete, no further optimization
DELETE_ROUND3 = {
    "g1000_office_desk__stand_keyboard_pad__occlusion_state_002",
    "g1000_warehouse_shelf__box_sack_pallet__2_to_1_001",
    "g1000_warehouse_shelf__box_sack_pallet__3_to_2_001",
    "g1000_warehouse_shelf__box_tote_drum__2_to_1_001",
    "g1000_warehouse_shelf__box_tote_drum__2_to_3_001",
    "g1000_warehouse_shelf__box_tote_drum__3_to_2_001",
}


def load_jsonl(path: Path) -> list[dict]:
    return [json.loads(line) for line in path.open(encoding="utf-8") if line.strip()]


def assets(batch_root: Path, record: dict) -> dict:
    def resolve(value: str) -> str:
        path = Path(str(value))
        return str(path if path.is_absolute() else batch_root / path)

    return {
        "source": resolve(record["source"]["image"]),
        "target": resolve(record["target"]["image"]),
        "mask": resolve(record["target"].get("mask", "")),
    }


def build() -> dict:
    decisions = {str(r["pair_id"]): r for r in load_jsonl(ROUND2_DECISIONS)}
    round2_manifest = {str(r["pair_id"]): r for r in load_jsonl(ROUND2_MANIFEST)}
    round3_manifest = {str(r["pair_id"]): r for r in load_jsonl(ROUND3_MANIFEST)}
    original = {str(r["pair_id"]): r for r in load_jsonl(ORIGINAL_MANIFEST)}

    rows: list[dict] = []
    for pid in sorted(decisions):
        base = decisions[pid]
        if base["decision"] == "delete":
            final = "deleted_round2"
        elif pid in DELETE_ROUND3:
            final = "deleted_round3"
        elif pid in round3_manifest:
            final = "pass_round3"
        else:
            final = "pass_round2"
        record = round3_manifest.get(pid) or round2_manifest.get(pid) or original.get(pid, {})
        batch_root = ROUND3_SANDBOX if pid in round3_manifest else ROUND2_SANDBOX
        path = assets(batch_root, record) if record else {}
        rows.append({
            "pair_id": pid,
            "source_scene_id": base["source_scene_id"],
            "subtype": base.get("subtype", ""),
            "intervention": base.get("intervention", ""),
            "round2_decision": base["decision"],
            "final_decision": final,
            "round3_issue": base.get("audit_issue", ""),
            "source_measure": base.get("source_measure", ""),
            "target_measure": base.get("target_measure", ""),
            "asset_batch": "round3" if pid in round3_manifest else "round2",
            "source_image": path.get("source", ""),
            "target_image": path.get("target", ""),
            "mask_image": path.get("mask", ""),
        })

    replace_rows = [r for r in rows if r["final_decision"].startswith("pass")]
    deleted_rows = [r for r in rows if r["final_decision"].startswith("deleted")]
    pass_round2 = [r for r in rows if r["final_decision"] == "pass_round2"]
    pass_round3 = [r for r in rows if r["final_decision"] == "pass_round3"]

    status_path = ROOT / "repair_status_20260928_final.jsonl"
    status_path.write_text("".join(json.dumps(r, ensure_ascii=False) + "\n" for r in rows), encoding="utf-8")

    replacement_path = ROOT / "final_replacement_20260928.jsonl"
    with replacement_path.open("w", encoding="utf-8") as handle:
        for row in replace_rows:
            record = round3_manifest.get(row["pair_id"]) or round2_manifest[row["pair_id"]]
            clean = json.loads(json.dumps(record))
            clean["source"]["image"] = row["source_image"]
            clean["target"]["image"] = row["target_image"]
            if row["mask_image"]:
                clean["target"]["mask"] = row["mask_image"]
            clean["replacement"] = {
                "replaced_at": "2026-09-28",
                "asset_batch": row["asset_batch"],
                "final_decision": row["final_decision"],
                "source_scene_id": row["source_scene_id"],
            }
            handle.write(json.dumps(clean, ensure_ascii=False) + "\n")

    summary = {
        "audit_date": "2026-09-28",
        "audited_pairs": len(rows),
        "pass_round2": len(pass_round2),
        "pass_round3": len(pass_round3),
        "to_replace_into_1000": len(replace_rows),
        "deleted_round2": sum(1 for r in deleted_rows if r["final_decision"] == "deleted_round2"),
        "deleted_round3": sum(1 for r in deleted_rows if r["final_decision"] == "deleted_round3"),
        "deleted_total": len(deleted_rows),
        "deleted_pair_ids": [r["pair_id"] for r in deleted_rows],
        "untouched_1000_pairs": len(original) - len(rows),
        "final_dataset_size_if_applied": (len(original) - len(rows)) + len(replace_rows),
        "replacement_jsonl": str(replacement_path),
        "status_jsonl": str(status_path),
    }
    (ROOT / "final_replacement_20260928_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    xlsx_path = ROOT / "repair_status_20260928_final.xlsx"
    build_workbook(rows, replace_rows, deleted_rows, summary).save(xlsx_path)
    summary["xlsx"] = str(xlsx_path)
    return summary


def build_workbook(rows, replace_rows, deleted_rows, summary) -> Workbook:
    head_fill = PatternFill("solid", fgColor="305496")
    head_font = Font(color="FFFFFF", bold=True)
    workbook = Workbook()

    def add_sheet(title: str, headers: list[str], data: list[list], widths: list[int]):
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for index in range(1, len(headers) + 1):
            cell = sheet.cell(row=1, column=index)
            cell.fill = head_fill
            cell.font = head_font
            sheet.column_dimensions[get_column_letter(index)].width = widths[index - 1]
        for row in data:
            sheet.append(row)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.freeze_panes = "A2"

    overview = workbook.active
    overview.title = "总览统计"
    overview.append(["项目", "数量"])
    for cell in overview[1]:
        cell.fill = head_fill
        cell.font = head_font
    for key, label in [
        ("audited_pairs", "本批审计 pair 总数"),
        ("pass_round2", "round2 通过（沿用 round2 资产）"),
        ("pass_round3", "round3 重生成后通过"),
        ("to_replace_into_1000", "需要替换进 1000 的 pair 数"),
        ("deleted_round2", "round2 审计删除"),
        ("deleted_round3", "round3 审计删除"),
        ("deleted_total", "删除合计"),
        ("untouched_1000_pairs", "1000 中未动、保持原样的 pair"),
        ("final_dataset_size_if_applied", "执行替换后的数据集规模"),
    ]:
        overview.append([label, summary[key]])
    overview.column_dimensions["A"].width = 38
    overview.column_dimensions["B"].width = 20

    add_sheet(
        "替换清单",
        ["pair_id", "source_scene_id", "subtype", "asset_batch", "final_decision", "source_image", "target_image", "mask_image"],
        [[r["pair_id"], r["source_scene_id"], r["subtype"], r["asset_batch"], r["final_decision"],
          r["source_image"], r["target_image"], r["mask_image"]] for r in replace_rows],
        [46, 40, 16, 12, 16, 60, 60, 60],
    )
    add_sheet(
        "删除清单",
        ["pair_id", "source_scene_id", "subtype", "final_decision", "原因"],
        [[r["pair_id"], r["source_scene_id"], r["subtype"], r["final_decision"], r["round3_issue"]]
         for r in deleted_rows],
        [46, 40, 16, 16, 60],
    )
    add_sheet(
        "全部明细",
        ["pair_id", "round2_decision", "final_decision", "asset_batch", "问题", "source 措施", "target 措施"],
        [[r["pair_id"], r["round2_decision"], r["final_decision"], r["asset_batch"],
          r["round3_issue"], r["source_measure"], r["target_measure"]] for r in rows],
        [46, 22, 18, 12, 50, 60, 60],
    )
    return workbook


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.parse_args()
    print(json.dumps(build(), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
