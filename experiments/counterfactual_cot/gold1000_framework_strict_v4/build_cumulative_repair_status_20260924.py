#!/usr/bin/env python3
"""Build one authoritative, one-row-per-pair repair status ledger.

The ledger starts from the 2026-09-23 Excel repair plan and overlays the
v3 candidate batch, the first follow-up repair batch, and the current strict
atomic round-two COT package. It never modifies the original 1000-pair data or
any repair asset.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path("/group/40016/niightli/BAGEL/experiments/counterfactual_cot/gold1000_framework_strict_v4")
DATA = ROOT / "full_1000_20260918"
ORIGINAL_MANIFEST = DATA / "gold1000.generated.jsonl"
INITIAL_XLSX = ROOT / "repair_plan_20260923.xlsx"
INITIAL_JSONL = ROOT / "repair_plan_20260923.jsonl"
V3_ROOT = ROOT / "repair_batch_20260923_v3"
V3_MANIFEST = V3_ROOT / "manifest.jsonl"
V1_ROOT = ROOT / "repair_batch_20240924_v1"
V1_MANIFEST = V1_ROOT / "manifest.jsonl"
ROUND2_DIRECT_PLAN = ROOT / "detailed_repair_plan_20260924_round2.jsonl"
ROUND2_ROOT = ROOT / "repair_round2_strict_atomic_20260924_v4"
ROUND2_MANIFEST = ROUND2_ROOT / "revised_cot_manifest.jsonl"
ROUND2_REGISTRY = ROUND2_ROOT / "revised_source_registry.jsonl"
ROUND2_TRACKING = ROUND2_ROOT / "round2_cot_tracking.jsonl"


def load_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.open(encoding="utf-8") if line.strip()]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def sha256_json(value: Any) -> str:
    return hashlib.sha256(
        json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()


def path_from_manifest(batch_root: Path, value: Any) -> str | None:
    if not value:
        return None
    path = Path(str(value))
    return str(path if path.is_absolute() else batch_root / path)


def file_status(path: str | None) -> str:
    if not path:
        return "missing_path"
    return "present" if Path(path).is_file() else "missing"


def read_initial_plan(path: Path) -> dict[str, dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["修补计划"]
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        raise RuntimeError("initial Excel repair plan is empty")
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    result: dict[str, dict[str, Any]] = {}
    for excel_row, values in enumerate(rows[1:], start=2):
        if not any(value is not None for value in values):
            continue
        row = {
            headers[index]: values[index] if index < len(values) else None
            for index in range(len(headers))
            if headers[index]
        }
        pair_id = str(row.get("pair_id") or "").strip()
        if not pair_id:
            raise RuntimeError(f"initial Excel row {excel_row} has no pair_id")
        if pair_id in result:
            raise RuntimeError(f"duplicate initial Excel pair_id: {pair_id}")
        row["_excel_row"] = excel_row
        result[pair_id] = row
    return result


def repair_role(row: dict[str, Any]) -> str | None:
    return (row.get("repair") or {}).get("repair_role")


def repair_scope(row: dict[str, Any]) -> str | None:
    return (row.get("repair") or {}).get("repair_scope")


def manifest_assets(row: dict[str, Any], batch_root: Path) -> dict[str, Any]:
    source_image = path_from_manifest(batch_root, (row.get("source") or {}).get("image"))
    target_image = path_from_manifest(batch_root, (row.get("target") or {}).get("image"))
    mask_image = path_from_manifest(batch_root, (row.get("target") or {}).get("mask"))
    return {
        "source_path": source_image,
        "target_path": target_image,
        "mask_path": mask_image,
        "source_file_status": file_status(source_image),
        "target_file_status": file_status(target_image),
        "mask_file_status": file_status(mask_image),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--label",
        default="pre_generation",
        help="snapshot label used in output filenames, e.g. pre_generation or post_generation",
    )
    parser.add_argument(
        "--round2-sandbox",
        type=Path,
        default=None,
        help="completed or in-progress round-two sandbox; use this to overlay generated manifest/ledger/report",
    )
    args = parser.parse_args()
    label = str(args.label).strip().replace(" ", "_")
    round2_sandbox = args.round2_sandbox.resolve() if args.round2_sandbox else None
    if not label:
        raise RuntimeError("label must not be empty")

    original = load_jsonl(ORIGINAL_MANIFEST)
    original_by_id = {str(row["pair_id"]): row for row in original}
    if len(original) != 1000 or len(original_by_id) != 1000:
        raise RuntimeError(f"expected 1000 unique original pairs, got rows={len(original)} unique={len(original_by_id)}")

    initial_by_id = read_initial_plan(INITIAL_XLSX)
    initial_jsonl = load_jsonl(INITIAL_JSONL)
    initial_jsonl_ids = {str(row["pair_id"]) for row in initial_jsonl}
    initial_ids = set(initial_by_id)
    if initial_ids != initial_jsonl_ids:
        raise RuntimeError(
            f"Excel/JSONL initial plan mismatch: Excel-only={sorted(initial_ids - initial_jsonl_ids)}, "
            f"JSONL-only={sorted(initial_jsonl_ids - initial_ids)}"
        )
    if not initial_ids <= set(original_by_id):
        raise RuntimeError("initial plan contains pair IDs absent from original manifest")

    v3 = load_jsonl(V3_MANIFEST)
    v1 = load_jsonl(V1_MANIFEST)
    v3_by_id = {str(row["pair_id"]): row for row in v3}
    v1_by_id = {str(row["pair_id"]): row for row in v1}
    if len(v3_by_id) != len(v3) or len(v1_by_id) != len(v1):
        raise RuntimeError("duplicate pair_id in v3 or v1 manifest")

    round2_direct_rows = load_jsonl(ROUND2_DIRECT_PLAN)
    round2_direct_by_id = {str(row["pair_id"]): row for row in round2_direct_rows}
    round2_manifest = load_jsonl(ROUND2_MANIFEST)
    round2_registry = load_jsonl(ROUND2_REGISTRY)
    round2_tracking = load_jsonl(ROUND2_TRACKING)
    round2_by_id = {str(row["pair_id"]): row for row in round2_manifest}
    round2_tracking_by_id = {str(row["pair_id"]): row for row in round2_tracking}
    if len(round2_by_id) != len(round2_manifest):
        raise RuntimeError("duplicate pair_id in strict round-two manifest")
    if len(round2_registry) != 30:
        raise RuntimeError(f"expected 30 revised source variants, got {len(round2_registry)}")
    if set(str(row["scene_id"]) for row in round2_registry) - {str(row["source_scene_id"]) for row in round2_manifest}:
        raise RuntimeError("revised source registry contains scene IDs absent from round-two manifest")
    if set(round2_direct_by_id) - set(round2_by_id):
        raise RuntimeError("round-two direct problem is absent from strict COT manifest")
    if set(round2_tracking_by_id) != set(round2_by_id):
        raise RuntimeError("strict round-two manifest/tracking pair sets differ")

    v3_ids = set(v3_by_id)
    v1_ids = set(v1_by_id)
    round2_ids = set(round2_by_id)
    round2_direct_ids = set(round2_direct_by_id)
    round2_spillover_ids = round2_ids - round2_direct_ids
    if len(round2_direct_ids) != 37 or len(round2_spillover_ids) != 12:
        raise RuntimeError(
            f"expected strict round-two 37 direct + 12 spillover, got {len(round2_direct_ids)} + {len(round2_spillover_ids)}"
        )
    if not round2_ids <= v1_ids:
        raise RuntimeError(f"strict round-two contains pair(s) absent from v1: {sorted(round2_ids - v1_ids)}")

    round2_generated_by_id: dict[str, dict[str, Any]] = {}
    round2_generation_report: dict[str, Any] = {}
    round2_ledger_events: list[dict[str, Any]] = []
    if round2_sandbox is not None:
        if not round2_sandbox.is_dir():
            raise RuntimeError(f"round-two sandbox does not exist: {round2_sandbox}")
        generated_manifest_path = round2_sandbox / "manifest.jsonl"
        if generated_manifest_path.is_file():
            generated_rows = load_jsonl(generated_manifest_path)
            round2_generated_by_id = {str(row["pair_id"]): row for row in generated_rows}
            if len(round2_generated_by_id) != len(generated_rows):
                raise RuntimeError("duplicate pair_id in round-two generated manifest")
            unknown_generated = sorted(set(round2_generated_by_id) - round2_ids)
            if unknown_generated:
                raise RuntimeError(f"generated round-two manifest contains unknown pair IDs: {unknown_generated[:10]}")
        report_path = round2_sandbox / "repair_report.json"
        if report_path.is_file():
            round2_generation_report = json.loads(report_path.read_text(encoding="utf-8"))
        ledger_path = round2_sandbox / "repair_ledger.jsonl"
        if ledger_path.is_file():
            round2_ledger_events = load_jsonl(ledger_path)

    round2_target_success_ids = {
        str(event["pair_id"])
        for event in round2_ledger_events
        if event.get("stage") == "target" and event.get("status") == "success"
    }
    round2_target_failed_ids = {
        str(event["pair_id"])
        for event in round2_ledger_events
        if event.get("stage") == "target" and event.get("status") in {"failed", "skipped_source_failed"}
    }
    round2_source_success_scene_ids = {
        str(event["scene_id"])
        for event in round2_ledger_events
        if event.get("stage") == "source" and event.get("status") == "success"
    }
    round2_report_summary = round2_generation_report.get("summary", {})
    round2_generation_observed = round2_sandbox is not None

    v3_candidate_ids = v3_ids - v1_ids
    v3_overlap_ids = v3_ids & v1_ids
    v1_default_after_round2_ids = v1_ids - round2_ids
    baseline_ids = set(original_by_id) - (v3_ids | v1_ids)
    current_partition = {
        "round2_pending": round2_ids,
        "v1_default_reasonable": v1_default_after_round2_ids,
        "v3_candidate": v3_candidate_ids,
        "original_baseline": baseline_ids,
    }
    if set().union(*current_partition.values()) != set(original_by_id):
        raise RuntimeError("current authoritative partition does not cover all original pairs")
    if sum(len(values) for values in current_partition.values()) != 1000:
        raise RuntimeError("current authoritative partition is not disjoint")

    records: list[dict[str, Any]] = []
    for original_line, original_row in enumerate(original, start=1):
        pair_id = str(original_row["pair_id"])
        initial = initial_by_id.get(pair_id)
        v3_row = v3_by_id.get(pair_id)
        v1_row = v1_by_id.get(pair_id)
        r2_row = round2_by_id.get(pair_id)
        r2_tracking_row = round2_tracking_by_id.get(pair_id)
        r2_direct = round2_direct_by_id.get(pair_id)

        if v3_row:
            v3_classification = "reported_problem" if pair_id in initial_ids else "source_scene_spillover"
            v3_status = "no_problem"
            v3_replacement = "pending" if pair_id in v3_candidate_ids else "hold_current_batch_overlap"
        else:
            v3_classification = "not_in_v3"
            v3_status = "not_selected"
            v3_replacement = "not_applicable"

        if v1_row:
            v1_role = repair_role(v1_row) or ("reported_problem" if pair_id in initial_ids else "source_scene_spillover")
            v1_scope = repair_scope(v1_row)
            v1_status = "generated_complete_pending_audit"
            v1_replacement = "blocked_until_manual_audit"
        else:
            v1_role = "not_in_v1"
            v1_scope = None
            v1_status = "not_selected"
            v1_replacement = "not_applicable"

        if pair_id in round2_ids:
            if pair_id in round2_direct_ids:
                round2_classification = "reported_problem"
                round2_user_listed = True
                round2_issue = r2_direct.get("user_issue")
            else:
                round2_classification = "source_scene_spillover"
                round2_user_listed = False
                round2_issue = "未列出；因共享 revised source variant 改变，必须同步重做 target。"
            round2_scope = "source+target" if r2_row["repair_round2"]["source_regenerated"] else "target-only"
            generated_row = round2_generated_by_id.get(pair_id)
            if not round2_generation_observed:
                round2_status = "strict_cot_ready_pending_generation"
                current_state = "round2_pending_generation"
                current_asset_status = "revised_cot_ready_no_image"
                current_replacement_status = "blocked_until_round2_generation_and_audit"
                current_reason_suffix = "等待 source/target/mask 生图与人工审计。"
            elif generated_row is not None:
                generated_assets = manifest_assets(generated_row, round2_sandbox)
                generated_complete = all(
                    generated_assets.get(key) == "present"
                    for key in ("source_file_status", "target_file_status", "mask_file_status")
                )
                if generated_complete:
                    round2_status = "generated_complete_pending_audit"
                    current_state = "round2_generated_pending_audit"
                    current_asset_status = "source_target_mask_generated_unreviewed"
                    current_replacement_status = "blocked_until_round2_manual_audit"
                    current_reason_suffix = "source/target/mask 已生成，等待分页人工审计。"
                else:
                    round2_status = "generated_manifest_asset_missing"
                    current_state = "round2_generation_asset_error"
                    current_asset_status = "generated_manifest_missing_asset"
                    current_replacement_status = "blocked_until_round2_generation_repair"
                    current_reason_suffix = "生成 manifest 存在，但 source/target/mask 文件不完整。"
            elif pair_id in round2_target_success_ids:
                round2_status = "target_success_pending_final_manifest"
                current_state = "round2_target_success_pending_manifest"
                current_asset_status = "target_success_ledger_source_mask_unverified"
                current_replacement_status = "blocked_until_round2_generation_and_audit"
                current_reason_suffix = "ledger 已记录 target success，等待最终 manifest/report 后核验 source/target/mask。"
            elif pair_id in round2_target_failed_ids:
                round2_status = "generation_failed"
                current_state = "round2_generation_failed"
                current_asset_status = "generation_failed_no_complete_asset"
                current_replacement_status = "blocked_until_round2_regeneration"
                current_reason_suffix = "target 生成失败或因 source 失败被跳过。"
            elif round2_generation_report:
                round2_status = "generation_report_missing_pair"
                current_state = "round2_generation_missing_pair"
                current_asset_status = "generation_report_incomplete"
                current_replacement_status = "blocked_until_round2_regeneration"
                current_reason_suffix = "生成报告已结束，但该 pair 未进入成功 manifest。"
            else:
                round2_status = "generation_in_progress"
                current_state = "round2_generation_in_progress"
                current_asset_status = "generation_in_progress"
                current_replacement_status = "blocked_until_round2_generation_and_audit"
                current_reason_suffix = "生图进程仍在运行，等待 source/target/mask 完成。"
            current_need_modification = True
            current_authoritative_batch = "repair_batch_20260924_round2_strict_atomic_v4"
            current_usable_by_policy = False
            current_reason = (
                "用户第二轮直接列出的问题，已完成严格 atomic COT 修订；"
                if round2_user_listed
                else "未列出但属于共享 revised source spillover；"
            ) + current_reason_suffix
        elif pair_id in v1_ids:
            round2_classification = "default_reasonable"
            round2_user_listed = False
            round2_issue = "用户第二轮未列出，按规则默认合理。"
            round2_scope = "not_required"
            round2_status = "default_reasonable_after_round2"
            current_state = "v1_default_reasonable_pending_fusion"
            current_need_modification = False
            current_authoritative_batch = "repair_batch_20240924_v1"
            current_asset_status = "v1_generated_unreviewed"
            current_replacement_status = "candidate_pending_fusion_integrity_check"
            current_usable_by_policy = True
            current_reason = "第二轮未列出；保留 v1 已生成资产，按用户规则作为默认合理候选。"
        elif pair_id in v3_candidate_ids:
            round2_classification = "default_reasonable"
            round2_user_listed = False
            round2_issue = "用户第二轮未列出，按规则默认合理。"
            round2_scope = "not_required"
            round2_status = "default_reasonable_after_round2"
            current_state = "v3_candidate_pending_fusion"
            current_need_modification = False
            current_authoritative_batch = "repair_batch_20260923_v3"
            current_asset_status = "v3_generated_no_problem"
            current_replacement_status = "candidate_pending_fusion_integrity_check"
            current_usable_by_policy = True
            current_reason = "未进入后续问题批次；v3 已记录 no_problem，保留为替换候选。"
        else:
            round2_classification = "default_reasonable"
            round2_user_listed = False
            round2_issue = "各轮均未列出，按用户规则默认合理；保留原始 1000-pair 基线。"
            round2_scope = "not_required"
            round2_status = "default_reasonable_baseline"
            current_state = "original_baseline_default_reasonable"
            current_need_modification = False
            current_authoritative_batch = "full_1000_20260918"
            current_asset_status = "original_baseline_retained"
            current_replacement_status = "baseline_retained"
            current_usable_by_policy = True
            current_reason = "各轮均未列出，按用户规则默认合理；原始资产保留在 full_1000 基线中。"

        conflict_flags: list[str] = []
        if pair_id in v3_overlap_ids:
            conflict_flags.append("cross_batch_overlap_v3_v1")
        if pair_id in round2_ids:
            conflict_flags.append("round2_reaudit_or_shared_source_sync")
        if pair_id in round2_direct_ids and pair_id not in initial_ids:
            conflict_flags.append("newly_listed_after_initial_plan")
        if pair_id in v1_ids and pair_id not in v3_ids:
            conflict_flags.append("newly_listed_after_v3")
        if pair_id in v3_ids and pair_id not in initial_ids:
            conflict_flags.append("v3_shared_source_spillover_from_initial_plan")
        if pair_id in round2_spillover_ids:
            conflict_flags.append("shared_revised_source_spillover")
        if pair_id in {"g1000_living_room_play__n1_rattle__3_to_2_001", "g1000_warehouse_shelf__bag_quota__3_to_2_001"}:
            if pair_id in round2_ids:
                conflict_flags.append("canonical_entity_or_count_revision")

        v3_assets = manifest_assets(v3_row, V3_ROOT) if v3_row else {}
        v1_assets = manifest_assets(v1_row, V1_ROOT) if v1_row else {}
        strict_assets = manifest_assets(r2_row, ROUND2_ROOT) if r2_row else {}
        generated_row = round2_generated_by_id.get(pair_id)
        generated_assets = (
            manifest_assets(generated_row, round2_sandbox)
            if generated_row is not None and round2_sandbox is not None
            else {}
        )
        r2_meta = r2_row.get("repair_round2", {}) if r2_row else {}
        r2_intervention = (r2_row or {}).get("intervention")

        record = {
            "schema_version": "repair_cumulative_status_v1",
            "snapshot_date": "2026-09-24",
            "snapshot_label": label,
            "pair_id": pair_id,
            "original_manifest_line": original_line,
            "source_scene_id_original": str(original_row["source_scene_id"]),
            "subtype": original_row.get("subtype"),
            "initial_plan_listed": initial is not None,
            "initial_plan_classification": "reported_problem" if initial is not None else "default_reasonable",
            "initial_plan_row": initial.get("_excel_row") if initial else None,
            "initial_plan_scope": initial.get("修补范围") if initial else "not_listed",
            "initial_plan_user_source_marked": initial.get("用户标记source") if initial else None,
            "initial_plan_user_issue": initial.get("用户问题") if initial else None,
            "initial_plan_intervention_field": initial.get("intervention_field") if initial else None,
            "initial_plan_old_value": initial.get("old_value") if initial else None,
            "initial_plan_new_value": initial.get("new_value") if initial else None,
            "initial_plan_generation_status": initial.get("生成状态") if initial else "not_in_initial_plan",
            "initial_plan_fusion_status": initial.get("融合状态") if initial else "baseline_or_not_selected",
            "round1_v3_present": bool(v3_row),
            "round1_v3_classification": v3_classification,
            "round1_v3_review_status": v3_status,
            "round1_v3_repair_scope": repair_scope(v3_row) if v3_row else None,
            "round1_v3_candidate_for_replacement": pair_id in v3_candidate_ids,
            "round1_v3_replacement_status": v3_replacement,
            "round1_v3_source_regenerated": (v3_row or {}).get("repair", {}).get("source_regenerated") if v3_row else None,
            "round1_v3_source_path": v3_assets.get("source_path"),
            "round1_v3_target_path": v3_assets.get("target_path"),
            "round1_v3_mask_path": v3_assets.get("mask_path"),
            "round1_v3_source_file_status": v3_assets.get("source_file_status"),
            "round1_v3_target_file_status": v3_assets.get("target_file_status"),
            "round1_v3_mask_file_status": v3_assets.get("mask_file_status"),
            "round1_v1_present": bool(v1_row),
            "round1_v1_classification": v1_role,
            "round1_v1_repair_scope": v1_scope,
            "round1_v1_user_issue": (v1_row or {}).get("repair", {}).get("user_issue") if v1_row else None,
            "round1_v1_review_status": "problem_under_repair" if v1_row else "not_selected",
            "round1_v1_generation_status": v1_status,
            "round1_v1_replacement_status": v1_replacement,
            "round1_v1_source_regenerated": (v1_row or {}).get("repair", {}).get("source_regenerated") if v1_row else None,
            "round1_v1_source_path": v1_assets.get("source_path"),
            "round1_v1_target_path": v1_assets.get("target_path"),
            "round1_v1_mask_path": v1_assets.get("mask_path"),
            "round1_v1_source_file_status": v1_assets.get("source_file_status"),
            "round1_v1_target_file_status": v1_assets.get("target_file_status"),
            "round1_v1_mask_file_status": v1_assets.get("mask_file_status"),
            "round2_present": pair_id in round2_ids,
            "round2_user_listed": round2_user_listed,
            "round2_classification": round2_classification,
            "round2_scope": round2_scope,
            "round2_user_issue": round2_issue,
            "round2_cot_revision_status": r2_tracking_row.get("cot_revision_status") if r2_tracking_row else "not_required",
            "round2_source_regeneration_required": bool(r2_meta.get("source_regenerated")) if r2_row else False,
            "round2_target_regeneration_required": bool(r2_row),
            "round2_source_variant_key": r2_tracking_row.get("source_variant_key") if r2_tracking_row else None,
            "round2_source_scene_id_revised": r2_tracking_row.get("source_scene_id_revised") if r2_tracking_row else None,
            "round2_intervention": r2_intervention,
            "round2_atomic_changed_atom_count": r2_tracking_row.get("atomic_changed_atom_count") if r2_tracking_row else None,
            "round2_generation_status": round2_status if pair_id in round2_ids else "not_required",
            "round2_replacement_status": current_replacement_status if pair_id in round2_ids else "not_required",
            "round2_fusion_status": "not_ready" if pair_id in round2_ids else "not_required_or_baseline",
            "round2_expected_source_path": strict_assets.get("source_path"),
            "round2_expected_target_path": strict_assets.get("target_path"),
            "round2_expected_mask_path": strict_assets.get("mask_path"),
            "round2_sandbox": str(round2_sandbox) if round2_sandbox else None,
            "round2_generated_manifest_present": generated_row is not None,
            "round2_generated_source_path": generated_assets.get("source_path"),
            "round2_generated_target_path": generated_assets.get("target_path"),
            "round2_generated_mask_path": generated_assets.get("mask_path"),
            "round2_generated_source_file_status": generated_assets.get("source_file_status"),
            "round2_generated_target_file_status": generated_assets.get("target_file_status"),
            "round2_generated_mask_file_status": generated_assets.get("mask_file_status"),
            "round2_ledger_source_success": bool(r2_row and str(r2_row["source_scene_id"]) in round2_source_success_scene_ids),
            "round2_ledger_target_success": pair_id in round2_target_success_ids,
            "round2_ledger_target_failed": pair_id in round2_target_failed_ids,
            "current_need_modification": current_need_modification,
            "current_state": current_state,
            "current_authoritative_batch": current_authoritative_batch,
            "current_asset_status": current_asset_status,
            "current_usable_by_user_default_policy": current_usable_by_policy,
            "current_replacement_status": current_replacement_status,
            "current_fusion_status": "not_fused_or_pending_integrity_check" if current_state != "original_baseline_default_reasonable" else "baseline_retained",
            "current_reason": current_reason,
            "conflict_flags": conflict_flags,
            "provenance": {
                "initial_plan": str(INITIAL_XLSX),
                "v3_manifest": str(V3_MANIFEST) if v3_row else None,
                "v1_manifest": str(V1_MANIFEST) if v1_row else None,
                "round2_manifest": str(ROUND2_MANIFEST) if r2_row else None,
                "round2_generated_manifest": str(round2_sandbox / "manifest.jsonl") if generated_row is not None and round2_sandbox else None,
                "round2_ledger": str(round2_sandbox / "repair_ledger.jsonl") if round2_sandbox and round2_ledger_events else None,
                "round2_report": str(round2_sandbox / "repair_report.json") if round2_sandbox and round2_generation_report else None,
            },
        }
        records.append(record)

    records.sort(key=lambda row: row["original_manifest_line"])
    out_jsonl = ROOT / f"repair_cumulative_status_20260924_{label}.jsonl"
    out_summary = ROOT / f"repair_cumulative_status_20260924_{label}_summary.json"
    out_xlsx = ROOT / f"repair_cumulative_status_20260924_{label}.xlsx"
    out_jsonl.write_text(
        "\n".join(json.dumps(record, ensure_ascii=False, sort_keys=True) for record in records) + "\n",
        encoding="utf-8",
    )

    current_state_counts = Counter(row["current_state"] for row in records)
    current_need_counts = Counter(row["current_need_modification"] for row in records)
    round2_counts = Counter(row["round2_classification"] for row in records)
    initial_counts = Counter(row["initial_plan_classification"] for row in records)
    round2_state_counts = Counter(
        row["current_state"] for row in records if row["round2_present"]
    )
    round2_generation_status = (
        "not_started"
        if not round2_generation_observed
        else "completed"
        if round2_generation_report
        and round2_report_summary.get("failures") == 0
        and round2_report_summary.get("generated_records") == round2_report_summary.get("selected_targets")
        else "completed_with_failures"
        if round2_generation_report
        else "in_progress_or_report_pending"
    )
    summary = {
        "schema_version": "repair_cumulative_status_summary_v1",
        "snapshot_date": "2026-09-24",
        "snapshot_label": label,
        "authoritative_rule": "每个pair只有一行；current_state按 round2 strict atomic v4 > v1 default reasonable > v3 candidate > original baseline 的优先级确定。",
        "original_manifest": str(ORIGINAL_MANIFEST),
        "original_manifest_sha256": sha256_file(ORIGINAL_MANIFEST),
        "initial_plan_xlsx": str(INITIAL_XLSX),
        "initial_plan_xlsx_sha256": sha256_file(INITIAL_XLSX),
        "input_hashes": {
            "v3_manifest": sha256_file(V3_MANIFEST),
            "v1_manifest": sha256_file(V1_MANIFEST),
            "round2_direct_plan": sha256_file(ROUND2_DIRECT_PLAN),
            "round2_manifest": sha256_file(ROUND2_MANIFEST),
            "round2_registry": sha256_file(ROUND2_REGISTRY),
            "round2_tracking": sha256_file(ROUND2_TRACKING),
            "round2_sandbox_manifest": (
                sha256_file(round2_sandbox / "manifest.jsonl")
                if round2_sandbox and (round2_sandbox / "manifest.jsonl").is_file()
                else None
            ),
            "round2_sandbox_ledger": (
                sha256_file(round2_sandbox / "repair_ledger.jsonl")
                if round2_sandbox and (round2_sandbox / "repair_ledger.jsonl").is_file()
                else None
            ),
            "round2_sandbox_report": (
                sha256_file(round2_sandbox / "repair_report.json")
                if round2_sandbox and (round2_sandbox / "repair_report.json").is_file()
                else None
            ),
        },
        "total_pairs": len(records),
        "initial_plan": {
            "listed_problem_count": len(initial_ids),
            "unlisted_default_reasonable_count": 1000 - len(initial_ids),
            "scope_counts": dict(Counter(str(row.get("修补范围")) for row in initial_by_id.values())),
            "source_marked_counts": dict(Counter(str(row.get("用户标记source")) for row in initial_by_id.values())),
        },
        "round1_v3": {
            "record_count": len(v3_ids),
            "initial_direct_count": len(v3_ids & initial_ids),
            "source_spillover_count": len(v3_ids - initial_ids),
            "no_problem_count": len(v3_ids),
            "candidate_pending_count": len(v3_candidate_ids),
            "held_by_v1_overlap_count": len(v3_overlap_ids),
            "candidate_rule": "v3 - v1 = pending candidate; v3 ∩ v1 = hold_current_batch_overlap",
        },
        "round1_v1": {
            "record_count": len(v1_ids),
            "direct_problem_count": sum(repair_role(row) == "reported_problem" for row in v1),
            "source_spillover_count": sum(repair_role(row) == "source_scene_spillover" for row in v1),
            "outside_initial_plan_count": len(v1_ids - initial_ids),
            "outside_v3_count": len(v1_ids - v3_ids),
            "generated_complete_pending_audit_count": len(v1_ids),
        },
        "round2_current_strict_atomic": {
            "record_count": len(round2_ids),
            "direct_problem_count": len(round2_direct_ids),
            "source_spillover_count": len(round2_spillover_ids),
            "source_plus_target_count": sum(bool(row["repair_round2"]["source_regenerated"]) for row in round2_manifest),
            "target_only_count": sum(not bool(row["repair_round2"]["source_regenerated"]) for row in round2_manifest),
            "outside_initial_plan_count": len(round2_ids - initial_ids),
            "direct_outside_initial_plan_count": len(round2_direct_ids - initial_ids),
            "source_variant_count": len(round2_registry),
            "source_scene_id_count_in_manifest": len({row["source_scene_id"] for row in round2_manifest}),
            "cot_revision_status": "revised_structured_cot_ready",
            "generation_status": round2_generation_status,
            "generation_snapshot": {
                "sandbox": str(round2_sandbox) if round2_sandbox else None,
                "ledger_event_count": len(round2_ledger_events),
                "source_success_scene_count": len(round2_source_success_scene_ids),
                "target_success_count": len(round2_target_success_ids),
                "target_failed_count": len(round2_target_failed_ids),
                "generated_manifest_count": len(round2_generated_by_id),
                "report_summary": round2_report_summary or None,
            },
            "state_counts": dict(round2_state_counts),
            "replacement_status": "blocked_until_round2_manual_audit" if round2_generation_status == "completed" else "blocked_until_round2_generation_and_audit",
        },
        "active_scope_trend": {
            "initial_plan_total_scope": len(initial_ids),
            "round1_v1_total_scope": len(v1_ids),
            "round2_current_total_scope": len(round2_ids),
            "initial_plan_direct_user_listed": len(initial_ids),
            "round1_v1_direct_user_listed": sum(repair_role(row) == "reported_problem" for row in v1),
            "round2_direct_user_listed": len(round2_direct_ids),
            "interpretation": "总修补范围 235 -> 81 -> 49；直接用户问题 235 -> 65 -> 37。spillover 已单独计数，不与 direct 混淆。",
        },
        "current_authoritative_partition": {
            **dict(current_state_counts),
            "default_reasonable_total": sum(
                1 for row in records if not row["current_need_modification"]
            ),
            "needs_modification_total": sum(
                1 for row in records if row["current_need_modification"]
            ),
            "partition_is_disjoint": True,
            "partition_covers_all_1000": True,
        },
        "current_state_counts": dict(current_state_counts),
        "current_need_modification_counts": {str(key): value for key, value in current_need_counts.items()},
        "round2_classification_counts": dict(round2_counts),
        "initial_classification_counts": dict(initial_counts),
        "fusion_gate": {
            "current_directly_fusable_count": 0,
            "candidate_pending_fusion_integrity_check_count": len(v3_candidate_ids) + len(v1_default_after_round2_ids),
            "baseline_retained_count": len(baseline_ids),
            "blocked_pending_round2_count": sum(
                1 for row in records if row["current_need_modification"]
            ),
            "round2_generated_pending_audit_count": sum(
                row["current_state"] == "round2_generated_pending_audit" for row in records
            ),
            "round2_generation_failed_count": sum(
                row["current_state"] in {"round2_generation_failed", "round2_generation_asset_error", "round2_generation_missing_pair"}
                for row in records
            ),
            "note": "用户默认合理不等于已执行融合；候选资产仍需 fusion integrity check，round2 完成生图后仍需分页人工审计。",
        },
        "legacy_tracking_note": {
            "old_tracking_rows": 342,
            "old_tracking_unique_pairs": 268,
            "old_duplicate_pair_rows": 74,
            "canonical_replacement": str(out_jsonl),
            "reason": "旧文件按batch一pair多行；本文件一pair一行，避免v3/v1/r2状态混淆。",
        },
        "special_sets": {
            "round2_direct_problem_pair_ids": sorted(round2_direct_ids),
            "round2_source_spillover_pair_ids": sorted(round2_spillover_ids),
            "round2_default_reasonable_after_v1_pair_ids": sorted(v1_default_after_round2_ids),
            "v3_candidate_pair_ids": sorted(v3_candidate_ids),
            "v3_v1_overlap_pair_ids": sorted(v3_overlap_ids),
            "initial_unlisted_but_later_v1_pair_ids": sorted((v1_ids - initial_ids)),
            "initial_unlisted_but_current_round2_pair_ids": sorted(round2_ids - initial_ids),
        },
        "files": {
            "jsonl": str(out_jsonl),
            "xlsx": str(out_xlsx),
            "summary": str(out_summary),
        },
    }
    out_summary.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    headers = list(records[0].keys())
    workbook = Workbook()
    status_sheet = workbook.active
    status_sheet.title = "pair_status"
    status_sheet.append(headers)
    for record in records:
        status_sheet.append(
            [
                json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value
                for value in (record.get(header) for header in headers)
            ]
        )
    status_sheet.freeze_panes = "A2"
    status_sheet.auto_filter.ref = status_sheet.dimensions
    for index, header in enumerate(headers, start=1):
        status_sheet.column_dimensions[get_column_letter(index)].width = min(max(len(header) + 2, 12), 32)

    summary_sheet = workbook.create_sheet("统计")
    summary_sheet.append(["统计项", "数量/值"])
    summary_rows = [
        ("总pair数", 1000),
        ("初始repair_plan列出问题", len(initial_ids)),
        ("初始未列出默认合理", 1000 - len(initial_ids)),
        ("v3总记录", len(v3_ids)),
        ("v3 no_problem候选", len(v3_ids)),
        ("v3 pending候选", len(v3_candidate_ids)),
        ("v3与v1重叠暂缓", len(v3_overlap_ids)),
        ("v1总修补范围", len(v1_ids)),
        ("v1直接问题", sum(repair_role(row) == "reported_problem" for row in v1)),
        ("v1 shared-source spillover", sum(repair_role(row) == "source_scene_spillover" for row in v1)),
        ("第二轮直接问题", len(round2_direct_ids)),
        ("第二轮shared-source spillover", len(round2_spillover_ids)),
        ("第二轮当前总修补范围", len(round2_ids)),
        ("第二轮生成状态", round2_generation_status),
        ("第二轮source成功scene", len(round2_source_success_scene_ids)),
        ("第二轮target成功", len(round2_target_success_ids)),
        ("第二轮target失败", len(round2_target_failed_ids)),
        ("当前需要修改", sum(row["current_need_modification"] for row in records)),
        ("当前默认合理", sum(not row["current_need_modification"] for row in records)),
        ("当前可直接融合", 0),
        ("当前候选待fusion integrity check", len(v3_candidate_ids) + len(v1_default_after_round2_ids)),
        ("当前原始基线保留", len(baseline_ids)),
    ]
    for summary_row in summary_rows:
        summary_sheet.append(summary_row)
    summary_sheet.append([])
    summary_sheet.append(["当前权威分区", "数量"])
    for state, count in sorted(current_state_counts.items()):
        summary_sheet.append([state, count])
    summary_sheet.append([])
    summary_sheet.append(["规则", "说明"])
    summary_sheet.append(["一行一pair", "本表是后续替换判断的唯一累计视图；旧batch tracking保留为历史文件。"])
    summary_sheet.append(["当前需要修改", f"当前共有{sum(row['current_need_modification'] for row in records)}条；其中第二轮 direct={len(round2_direct_ids)}、spillover={len(round2_spillover_ids)}。"])
    summary_sheet.append(["默认合理", "未列出的pair按用户规则默认合理，但候选资产仍需融合完整性检查。"])
    summary_sheet.append(["可直接融合", "必须人工审计通过且完成source/target/mask一致性检查；当前为0。"])
    summary_sheet.append(["round2生成观察", f"sandbox={round2_sandbox or '未提供'}；source成功scene={len(round2_source_success_scene_ids)}；target成功={len(round2_target_success_ids)}；target失败={len(round2_target_failed_ids)}。"])
    summary_sheet.freeze_panes = "A2"
    summary_sheet.column_dimensions["A"].width = 42
    summary_sheet.column_dimensions["B"].width = 80

    conflict_sheet = workbook.create_sheet("冲突与状态规则")
    conflict_sheet.append(["项目", "结果"])
    conflict_sheet.append(["初始计划", "235 listed problem + 765 unlisted default reasonable"])
    conflict_sheet.append(["v3", "261 no_problem; 187 pending candidate; 74 held by v1 overlap"])
    conflict_sheet.append(["v1", "81 total = 65 direct problem + 16 source spillover; generated, pending audit"])
    conflict_sheet.append(["round2 strict atomic", f"{len(round2_ids)} total = {len(round2_direct_ids)} direct problem + {len(round2_spillover_ids)} source spillover; COT ready"])
    conflict_sheet.append(["round2 generation", f"{round2_generation_status}; source_success={len(round2_source_success_scene_ids)}, target_success={len(round2_target_success_ids)}, target_failed={len(round2_target_failed_ids)}"])
    conflict_sheet.append(["current unresolved", f"{sum(row['current_need_modification'] for row in records)}; all are blocked until round2 audit/regeneration as applicable"])
    conflict_sheet.append(["current default reasonable", f"{sum(not row['current_need_modification'] for row in records)}; includes candidate and retained baseline states"])
    conflict_sheet.append(["old round2 plan", "51 total / 14 spillover is historical; current authoritative package is strict atomic v4"])
    conflict_sheet.append(["fusion", "No pair is directly fusable in this snapshot"])
    conflict_sheet.append(["sync command", "Run this script with --round2-sandbox <sandbox> --label post_generation after repair_report.json exists"] )
    conflict_sheet.freeze_panes = "A2"
    conflict_sheet.column_dimensions["A"].width = 30
    conflict_sheet.column_dimensions["B"].width = 110

    workbook.save(out_xlsx)
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
