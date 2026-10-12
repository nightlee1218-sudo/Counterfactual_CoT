#!/usr/bin/env python3
"""Build a per-pair, COT-aware repair plan for human approval.

This is planning only: no API calls, image generation, manifest replacement,
or pair deletion. The output keeps original COT/intervention/preservation
unchanged and proposes generation-prompt additions that preserve the existing
single-atomic-change contract.
"""
from __future__ import annotations

import argparse
import copy
import json
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

COUNT_SUBTYPES = {"1_to_2", "2_to_1", "2_to_3", "3_to_2", "3_to_1", "1_to_3"}
RELATION_SUBTYPES = {"occlusion_state", "containment"}
MOVE_SUBTYPES = {"absolute_horizontal", "absolute_vertical"}
ATTRIBUTE_SUBTYPES = {"material", "pattern", "functional_state", "color"}

PROMPT_TAIL = (
    " Keep a native 1024x1024 square, full-bleed scene with the same camera, lighting, background and composition. "
    "Do not introduce a frame, border, blank margin, collage, card, poster, page, mockup, text, logo, watermark or UI."
)


def load_manifest(path: Path) -> dict[str, dict[str, Any]]:
    return {r["pair_id"]: r for r in (json.loads(line) for line in path.open(encoding="utf-8") if line.strip())}


def load_analysis(path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["修补分析"]
    headers = [c.value for c in ws[1]]
    return [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]


def row_value(row: dict[str, Any], *names: str, default: Any = "") -> Any:
    """Read compatible column aliases from the two analysis-XLSX versions."""
    for name in names:
        value = row.get(name)
        if value is not None and str(value).strip() != "":
            return value
    return default


def entity_map(pair: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {str(e.get("id")): e for e in pair.get("source", {}).get("cot", {}).get("entities", [])}


def entity_phrase(entity: dict[str, Any]) -> str:
    attrs = entity.get("attributes", {}) or {}
    values: list[str] = []
    for key in ("color", "material", "pattern"):
        value = str(attrs.get(key, "")).replace("_", " ").strip()
        if value and value.lower() not in {x.lower() for x in values}:
            values.append(value)
    category = str(entity.get("category", "object")).replace("_", " ")
    values.append(category)
    return " ".join(values)


def related_entity(pair: dict[str, Any], entity: dict[str, Any]) -> dict[str, Any] | None:
    emap = entity_map(pair)
    for relation in entity.get("relations", []) or []:
        target = relation.get("target") or relation.get("target_id")
        if target in emap:
            return emap[target]
    return None


def entity_summary(pair: dict[str, Any]) -> str:
    return "; ".join(
        f"{e.get('id')}={e.get('count')}x {entity_phrase(e)} @ {e.get('location')}"
        for e in pair.get("source", {}).get("cot", {}).get("entities", [])
    )


def issue_has(issue: str, *terms: str) -> bool:
    lower = issue.lower()
    return any(term.lower() in lower for term in terms)


def source_addendum(pair: dict[str, Any], issue: str) -> str:
    iv = pair.get("intervention", {})
    subtype = str(pair.get("subtype", ""))
    entity, related = _target_and_related(pair)
    obj = entity_phrase(entity)
    clauses = [
        "Treat the original Source COT above as authoritative; do not change its entity IDs, counts, categories, attributes, locations or relations.",
        "Show every COT-listed entity exactly at its COT count, with no duplicate row, duplicate pile, hidden extra instance or unlisted object.",
        "Use a complete native square scene, not a crop, cutout, frame or partial scene.",
    ]
    if subtype in COUNT_SUBTYPES or iv.get("field") == "count":
        count = entity.get("count")
        clauses.append(
            f"For entity_id={iv.get('entity_id')}, show exactly {count} individual visible {obj} instances in one clearly bounded group/one pile when the COT describes a group. "
            "The number counts individual objects or individual folded garment layers, never piles, rows, shelves or containers."
        )
        clauses.append("Keep other same-category entities out of the target area so the old_value is unambiguous; leave free surface in the intended change direction.")
    if subtype in RELATION_SUBTYPES:
        if related:
            clauses.append(f"The only relation reference for the target is the existing {entity_phrase(related)}; do not invent another occluder or container.")
        old = str(iv.get("old_value", ""))
        if subtype == "occlusion_state":
            if old == "unoccluded" or old.startswith("unoccluded"):
                clauses.append("Initial state must be visibly unoccluded: a clear gap, no overlap, and the complete target outline visible.")
            else:
                clauses.append("Initial state must show a clear, judgeable partial overlap with both existing objects visible and roughly one-third of the target outline covered.")
        else:
            clauses.append("Initial containment must be physically judgeable: the object and container must both remain visible, and the container must be larger than the contained object.")
    if subtype in MOVE_SUBTYPES:
        clauses.append(f"Leave clear physical empty space in the requested destination {iv.get('new_value')} for the entire {obj}; do not place another object in that destination.")
    if issue_has(issue, "两行", "两沓", "太多", "太杂", "满满", "复杂", "四个", "三层"):
        clauses.append("Simplify the composition to only the COT-listed objects and one unambiguous arrangement; never create repeated rows or repeated levels of the same objects.")
    if issue_has(issue, "空间不足", "拥挤", "妨碍", "挤", "变短", "位置变化"):
        clauses.append("Reserve a visible empty area around the target operation so no unrelated object needs to move, resize, shorten, rotate or be occluded.")
    if issue_has(issue, "黑框", "白框", "抠图"):
        clauses.append("Fill the whole canvas with the intended scene and remove any black/white border, cutout artifact or blank wedge.")
    if issue_has(issue, "不明显", "不明确", "没有遮盖", "遮盖关系"):
        clauses.append("Make the initial relation visually unmistakable at a glance; do not rely on a tiny or ambiguous overlap.")
    return " ".join(clauses)


def target_addendum(pair: dict[str, Any], issue: str) -> str:
    iv = pair.get("intervention", {})
    subtype = str(pair.get("subtype", ""))
    entity, related = _target_and_related(pair)
    obj = entity_phrase(entity)
    eid = iv.get("entity_id")
    old, new = iv.get("old_value"), iv.get("new_value")
    clauses: list[str] = [f"This is the only allowed atomic change: entity_id={eid}, field={iv.get('field')}, {old} -> {new}."]
    if subtype in COUNT_SUBTYPES or iv.get("field") == "count":
        if isinstance(old, int) and isinstance(new, int) and new > old:
            clauses.append(
                f"Add exactly {new - old} individual {obj} instance(s) inside the same existing target group/one pile, directly on the existing group in the same row and same area. "
                f"Do not create a second pile, second row, second shelf, new area or duplicate the whole group. Final individual count is exactly {new}, not {new + 1} and not {new * 2}."
            )
        elif isinstance(old, int) and isinstance(new, int):
            clauses.append(
                f"Remove exactly {old - new} individual {obj} instance(s) from the same existing target group/one pile. Survivors keep position, size and distinctive details. Final individual count is exactly {new}; do not remove the whole group or replace survivors."
            )
        clauses.append("No unrelated object may move, resize, shorten, rotate, disappear or change material/state.")
    elif subtype in MOVE_SUBTYPES:
        clauses.append(
            f"Move the existing {obj} from {old} to {new} without changing its count, size, orientation or distinctive details. Clear the old position; do not move or alter any unrelated object."
        )
    elif subtype == "occlusion_state":
        ref = entity_phrase(related) if related else "the existing reference object"
        clauses.append(
            f"Use the existing {ref}; do not create a second one. The requested relation must be obvious: both objects remain present, and the target object's visible outline changes exactly as required by the target COT."
        )
    elif subtype == "containment":
        ref = entity_phrase(related) if related else "the existing container"
        clauses.append(
            f"Use the existing {ref}; do not add or delete a container/object. Change only the relation, keep the container position, shape, size and count unchanged, and keep the object visible enough to verify the relation."
        )
    elif subtype in ATTRIBUTE_SUBTYPES or str(iv.get("field", "")).startswith("attributes."):
        clauses.append(
            f"Make the requested {iv.get('field')} change clearly visible on the existing {obj}, while keeping count, position, silhouette, identity and every other attribute unchanged. A subtle or ambiguous change is a failure."
        )
    elif subtype in {"affordance_divergent_replace", "shape_divergent_replace", "same_affordance_replace"}:
        clauses.append(
            f"Remove the old {old} completely and replace the same single instance at the same position, count, size and orientation with an unmistakable {new} silhouette. Do not leave the old object or create a second object."
        )
    if issue_has(issue, "蜡", "流蜡"):
        clauses.append("The surviving candle must retain its original melted wax, wax drips, stains, burn marks and other distinctive details; never redraw it as a clean new candle.")
    if issue_has(issue, "旋转", "方向", "放倒", "切割", "变短", "挤", "位置"):
        clauses.append("Preserve unrelated geometry exactly; no rotation, shortening, cutting, displacement or restyling of unrelated objects.")
    return " ".join(clauses)


def _target_and_related(pair: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any] | None]:
    entity, _ = _entity_and_map(pair)
    return entity, related_entity(pair, entity)


def _entity_and_map(pair: dict[str, Any]) -> tuple[dict[str, Any], dict[str, dict[str, Any]]]:
    iv = pair.get("intervention", {})
    entities = pair.get("source", {}).get("cot", {}).get("entities", [])
    emap = {str(e.get("id")): e for e in entities}
    return emap.get(str(iv.get("entity_id")), {}), emap


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--analysis-xlsx", required=True, type=Path)
    ap.add_argument("--manifest", required=True, type=Path)
    ap.add_argument("--output-xlsx", required=True, type=Path)
    ap.add_argument("--output-jsonl", required=True, type=Path)
    ap.add_argument("--output-json", required=True, type=Path)
    args = ap.parse_args()

    rows = load_analysis(args.analysis_xlsx)
    manifest = load_manifest(args.manifest)
    scene_members: dict[str, list[str]] = defaultdict(list)
    for pid, record in manifest.items():
        scene_members[str(record["source_scene_id"])].append(pid)

    dedup: dict[str, dict[str, Any]] = {}
    for row in rows:
        pid = row_value(row, "pair_id_manifest", "pair_id_raw", "pair_id")
        if not pid or pid not in manifest:
            continue
        if pid not in dedup or row_value(row, "用户明确source重生成", "用户source标记") == "是":
            dedup[pid] = row

    plans: list[dict[str, Any]] = []
    for pid, audit_row in sorted(dedup.items()):
        pair = manifest[pid]
        issue = str(row_value(audit_row, "用户问题", "本次审计问题") or "")
        scope = str(row_value(audit_row, "repair_scope", "建议修补范围", default="target-only") or "target-only")
        entity, _ = _entity_and_map(pair)
        source_ids = sorted(scene_members[str(pair["source_scene_id"])])
        source_text = source_addendum(pair, issue) if scope == "source+target" else "不修改Source；使用最终manifest中的现有Source作为Target编辑输入。"
        target_text = target_addendum(pair, issue)
        original_source_prompt = pair.get("source", {}).get("prompt", "")
        original_target_prompt = pair.get("target", {}).get("prompt", "")
        proposed_source_prompt = original_source_prompt + "\n\nREPAIR SOURCE ADDENDUM:\n" + source_text if scope == "source+target" else ""
        proposed_target_prompt = original_target_prompt + "\n\nREPAIR TARGET ADDENDUM:\n" + target_text
        ambiguity = "高" if (isinstance(entity.get("count"), int) and entity.get("count", 0) > 1 and pair.get("subtype") in COUNT_SUBTYPES | RELATION_SUBTYPES) else "低"
        if pair.get("subtype") in COUNT_SUBTYPES and isinstance(entity.get("count"), int) and entity.get("count", 0) > 1:
            ambiguity_reason = f"{pair['intervention'].get('entity_id')} 是组级entity，当前count={entity.get('count')}；后续必须按individual实例/同一摞内部层数解释，不能把entity_id当作第一个独立物体。"
        else:
            ambiguity_reason = "未发现明确组级计数歧义；仍需按source/target COT和图片核对。"
        plans.append({
            "pair_id": pid, "source_scene_id": pair["source_scene_id"], "scene_pair_count": len(source_ids), "scene_pair_ids": source_ids,
            "subtype": pair.get("subtype"), "intervention": copy.deepcopy(pair.get("intervention")), "preservation": copy.deepcopy(pair.get("preservation")),
            "entity_summary": entity_summary(pair), "entity_id": pair.get("intervention", {}).get("entity_id"), "entity_group_count": entity.get("count"),
            "entity_count_ambiguity": ambiguity, "entity_ambiguity_reason": ambiguity_reason, "user_issue": issue,
            "user_source_marked": row_value(audit_row, "用户明确source重生成", "用户source标记") == "是", "repair_scope": scope,
            "source_measure": source_text, "target_measure": target_text, "original_source_prompt": original_source_prompt,
            "original_target_prompt": original_target_prompt, "proposed_source_prompt": proposed_source_prompt,
            "proposed_target_prompt": proposed_target_prompt, "source_cot_text": pair.get("source", {}).get("cot_text"),
            "target_cot_text": pair.get("target", {}).get("cot_text"), "approval_status": "待人工审核",
            "generation_status": "未生成", "merge_status": "未融合",
        })

    args.output_jsonl.parent.mkdir(parents=True, exist_ok=True)
    with args.output_jsonl.open("w", encoding="utf-8") as f:
        for plan in plans:
            f.write(json.dumps(plan, ensure_ascii=False, sort_keys=True) + "\n")
    args.output_json.write_text(json.dumps({"schema_version": "counterfactual_detailed_repair_plan_v1", "manifest": str(args.manifest), "pair_count": len(plans), "plans": plans}, ensure_ascii=False, indent=2), encoding="utf-8")

    wb = Workbook(); ws = wb.active; ws.title = "逐Pair详细修补计划"
    headers = [
        "序号", "pair_id", "source_scene_id", "scene_pair_count", "scene_pair_ids", "subtype", "intervention_field", "old_value", "new_value", "entity_id", "entity_group_count", "entity_count_ambiguity", "entity_ambiguity_reason", "entity_summary", "用户问题", "用户source标记", "repair_scope", "Source具体修改措施", "Target具体修改措施", "单原子/must_preserve", "验收检查", "原SourcePrompt", "建议SourcePrompt", "原TargetPrompt", "建议TargetPrompt", "source_cot_text", "target_cot_text", "审批状态", "生成状态", "融合状态"
    ]
    ws.append(headers)
    for i, plan in enumerate(plans, 1):
        iv = plan["intervention"] or {}
        ws.append([
            i, plan["pair_id"], plan["source_scene_id"], plan["scene_pair_count"], ";".join(plan["scene_pair_ids"]), plan["subtype"], iv.get("field"), iv.get("old_value"), iv.get("new_value"), plan["entity_id"], plan["entity_group_count"], plan["entity_count_ambiguity"], plan["entity_ambiguity_reason"], plan["entity_summary"], plan["user_issue"], "是" if plan["user_source_marked"] else "否", plan["repair_scope"], plan["source_measure"], plan["target_measure"], "仅允许 intervention.entity_id 的 intervention.field 变化；其他must_preserve字段不变。", "核对Source old状态/数量、Target new状态/数量、无新增/消失、无关物体原位、单原子性、mask局部性。", plan["original_source_prompt"], plan["proposed_source_prompt"], plan["original_target_prompt"], plan["proposed_target_prompt"], plan["source_cot_text"], plan["target_cot_text"], plan["approval_status"], plan["generation_status"], plan["merge_status"]
        ])
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F4E78"); cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = [8, 60, 42, 14, 100, 22, 28, 14, 14, 20, 14, 16, 72, 90, 75, 14, 16, 100, 120, 80, 90, 110, 130, 120, 140, 100, 100, 16, 16, 16]
    for i, width in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = width
    table = Table(displayName="DetailedRepairPlan20240924", ref=f"A1:AD{ws.max_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False); ws.add_table(table)

    stats = wb.create_sheet("统计")
    stats.append(["统计项", "数量"])
    stats.append(["唯一pair", len(plans)])
    stats.append(["source+target", sum(p["repair_scope"] == "source+target" for p in plans)])
    stats.append(["target-only", sum(p["repair_scope"] == "target-only" for p in plans)])
    stats.append(["涉及共享source的pair", sum(p["scene_pair_count"] > 1 for p in plans)])
    stats.append(["高entity计数歧义", sum(p["entity_count_ambiguity"] == "高" for p in plans)])
    stats.append([]); stats.append(["subtype", "数量"])
    for subtype, count in sorted(Counter(str(p["subtype"]) for p in plans).items()): stats.append([subtype, count])
    for cell in stats[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F4E78")
    stats.column_dimensions["A"].width = 35; stats.column_dimensions["B"].width = 15

    args.output_xlsx.parent.mkdir(parents=True, exist_ok=True); wb.save(args.output_xlsx)
    print(json.dumps({"output_xlsx": str(args.output_xlsx), "output_jsonl": str(args.output_jsonl), "output_json": str(args.output_json), "unique_pairs": len(plans), "source_target": sum(p["repair_scope"] == "source+target" for p in plans), "target_only": sum(p["repair_scope"] == "target-only" for p in plans), "shared_scene_pairs": sum(p["scene_pair_count"] > 1 for p in plans)}, ensure_ascii=False, indent=2))


if __name__ == "__main__": main()
