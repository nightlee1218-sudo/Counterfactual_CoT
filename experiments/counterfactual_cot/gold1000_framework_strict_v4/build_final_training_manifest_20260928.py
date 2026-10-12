#!/usr/bin/env python3
"""Build a READ-ONLY final training manifest (993 pairs) with resolved asset paths.

Never deletes, never moves, never merges. It only lists what survived all audits.

Outputs
  - final_training_manifest_20260928.jsonl   993 full records + training_assets block
  - final_training_manifest_20260928.xlsx    overview / manifest / issue flags
  - final_training_manifest_20260928_summary.json
"""
from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path("/group/40016/niightli/BAGEL/experiments/counterfactual_cot/gold1000_framework_strict_v4")
BASE_ROOT = ROOT / "full_1000_20260918" / "source_images"
BASE_MANIFEST = ROOT / "full_1000_20260918" / "gold1000.generated.jsonl"
REPLACEMENT = ROOT / "final_replacement_20260928_full.jsonl"
DELETIONS = ROOT / "repair_status_20260928_final.jsonl"

# 2026-09-28 final human adjudication from final_audit_page_20260928
#   1. scene conflict group: only ball_quota__3_to_2_001 is removed
#   2. quality-failed group: all 14 pass
#   3. merged-36 group: only glasses_charger_glass__occlusion_state_005 is removed
FINAL_EXCLUSIONS = {
    "g1000_living_room_play__ball_quota__3_to_2_001",
    "g1000_bedroom_nightstand__glasses_charger_glass__occlusion_state_005",
}


def load_jsonl(path: Path) -> list[dict]:
    return [json.loads(line) for line in path.open(encoding="utf-8") if line.strip()]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def resolve(value: str, batch_root: Path | None) -> Path:
    path = Path(str(value))
    if path.is_absolute():
        return path
    return (batch_root or BASE_ROOT) / path


def build() -> dict:
    base = {str(r["pair_id"]): r for r in load_jsonl(BASE_MANIFEST)}
    replacement = {str(r["pair_id"]): r for r in load_jsonl(REPLACEMENT)}
    deleted = {str(r["pair_id"]) for r in load_jsonl(DELETIONS)
               if str(r.get("final_decision", "")).startswith("deleted")}
    deleted |= FINAL_EXCLUSIONS

    batch_root = {
        "repair_batch_20260923_v3": ROOT / "repair_batch_20260923_v3",
        "repair_batch_20240924_v1": ROOT / "repair_batch_20240924_v1",
        "repair_batch_20260924_round2_strict_atomic_v4": ROOT / "repair_batch_20260924_round2_strict_atomic_v4",
        "repair_batch_20260928_round3": ROOT / "repair_batch_20260928_round3",
    }

    final: dict[str, dict] = {}
    meta: dict[str, dict] = {}
    for pid, record in base.items():
        if pid in deleted:
            continue
        final[pid] = ("baseline", record)
        meta[pid] = {"batch_root": None, "asset_batch": "baseline"}
    for pid, record in replacement.items():
        if pid in deleted:
            continue
        batch = str(record["replacement"]["asset_batch"])
        final[pid] = (batch, record)
        meta[pid] = {
            "batch_root": batch_root.get({"v3": "repair_batch_20260923_v3",
                                          "v1": "repair_batch_20240924_v1",
                                          "round2": "repair_batch_20260924_round2_strict_atomic_v4",
                                          "round3": "repair_batch_20260928_round3"}[batch]),
            "asset_batch": batch,
        }

    # scene source consistency
    scene_sources: dict[str, dict[str, list[str]]] = defaultdict(lambda: defaultdict(list))
    rows: list[dict] = []
    records_out: list[dict] = []
    missing: list[tuple[str, str]] = []
    no_mask: list[str] = []
    quality_fail: list[tuple[str, str, str]] = []

    for pid in sorted(final):
        batch, record = final[pid]
        root = meta[pid]["batch_root"]
        source = resolve(str(record["source"]["image"]), root)
        target = resolve(str(record["target"]["image"]), root)
        mask_value = str(record["target"].get("mask", ""))
        mask = resolve(mask_value, root) if mask_value else None

        if not source.is_file():
            missing.append((pid, "source"))
        if not target.is_file():
            missing.append((pid, "target"))
        if mask is None or not mask.is_file():
            no_mask.append(pid)

        scene_id = str(record.get("source_scene_id") or record.get("source", {}).get("source_scene_id") or "")
        if source.is_file():
            scene_sources[scene_id][sha256_file(source)].append(pid)

        checks = record.get("quality", {}).get("local_checks", {}) or {}
        if checks and not checks.get("passed", True):
            quality_fail.append((pid, batch, ",".join(checks.get("failures", []))))

        out = json.loads(json.dumps(record))
        out["source"]["image"] = str(source)
        out["target"]["image"] = str(target)
        if mask is not None:
            out["target"]["mask"] = str(mask)
        out["training_assets"] = {
            "pair_id": pid,
            "asset_batch": batch,
            "human_approved": batch != "baseline",
            "source_abs": str(source),
            "target_abs": str(target),
            "mask_abs": str(mask) if mask is not None else "",
            "source_sha256": sha256_file(source) if source.is_file() else "",
            "target_sha256": sha256_file(target) if target.is_file() else "",
            "mask_sha256": sha256_file(mask) if mask is not None and mask.is_file() else "",
        }
        records_out.append(out)
        rows.append({
            "pair_id": pid,
            "subtype": record.get("subtype", ""),
            "source_scene_id": scene_id,
            "asset_batch": batch,
            "human_approved": "是" if batch != "baseline" else "原始基线",
            "source_abs": str(source),
            "target_abs": str(target),
            "mask_abs": str(mask) if mask is not None else "",
            "target_sha256": out["training_assets"]["target_sha256"][:16],
        })

    scene_conflicts = {sid: {k: v for k, v in shas.items()}
                       for sid, shas in scene_sources.items() if len(shas) > 1}
    for row in rows:
        sid = row["source_scene_id"]
        conflict = sid in scene_conflicts
        row["scene_source_conflict"] = "是" if conflict else ""
        row["quality_local_passed"] = "否" if any(pid == row["pair_id"] for pid, _, _ in quality_fail) else "是"

    out_path = ROOT / "final_training_manifest_20260928.jsonl"
    with out_path.open("w", encoding="utf-8") as handle:
        for record in records_out:
            handle.write(json.dumps(record, ensure_ascii=False) + "\n")

    summary = {
        "created_at": "2026-09-28",
        "mode": "read_only_listing",
        "final_pair_count": len(records_out),
        "asset_batch_counts": dict(Counter(r["asset_batch"] for r in rows)),
        "deleted_pair_ids": sorted(deleted),
        "missing_files": missing,
        "missing_file_count": len(missing),
        "pairs_without_mask": no_mask,
        "quality_local_failed_count": len(quality_fail),
        "quality_local_failed": quality_fail,
        "scene_source_conflict_count": len(scene_conflicts),
        "scene_source_conflicts": {sid: {"pair_count": sum(len(v) for v in shas.values()),
                                         "variants": {k[:12]: v for k, v in shas.items()}}
                                   for sid, shas in scene_conflicts.items()},
        "manifest_jsonl": str(out_path),
    }
    (ROOT / "final_training_manifest_20260928_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    xlsx_path = ROOT / "final_training_manifest_20260928.xlsx"
    build_workbook(rows, summary).save(xlsx_path)
    summary["xlsx"] = str(xlsx_path)
    return summary


def build_workbook(rows, summary) -> Workbook:
    head_fill = PatternFill("solid", fgColor="305496")
    head_font = Font(color="FFFFFF", bold=True)
    warn_fill = PatternFill("solid", fgColor="FFC7CE")
    workbook = Workbook()

    overview = workbook.active
    overview.title = "总览"
    overview.append(["项目", "数量"])
    for cell in overview[1]:
        cell.fill = head_fill
        cell.font = head_font
    overview.append(["最终保留 pair 总数", summary["final_pair_count"]])
    for batch, count in sorted(summary["asset_batch_counts"].items()):
        overview.append([f"资产来源 {batch}", count])
    overview.append(["删除（不进入训练）", len(summary["deleted_pair_ids"])])
    overview.append(["文件缺失", summary["missing_file_count"]])
    overview.append(["无 mask", len(summary["pairs_without_mask"])])
    overview.append(["自动质量未通过", summary["quality_local_failed_count"]])
    overview.append(["同场景 source 冲突", summary["scene_source_conflict_count"]])
    overview.column_dimensions["A"].width = 30
    overview.column_dimensions["B"].width = 22

    sheet = workbook.create_sheet("训练清单")
    headers = ["pair_id", "subtype", "source_scene_id", "asset_batch", "是否人审通过",
               "同场景source冲突", "自动质量通过", "source_abs", "target_abs", "mask_abs", "target_sha256"]
    widths = [46, 18, 40, 12, 14, 14, 12, 60, 60, 60, 20]
    sheet.append(headers)
    for index in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=index)
        cell.fill = head_fill
        cell.font = head_font
        sheet.column_dimensions[get_column_letter(index)].width = widths[index - 1]
    for row in rows:
        sheet.append([row[h] for h in ["pair_id", "subtype", "source_scene_id", "asset_batch", "human_approved",
                                       "scene_source_conflict", "quality_local_passed", "source_abs",
                                       "target_abs", "mask_abs", "target_sha256"]])
    for excel_row in sheet.iter_rows(min_row=2):
        for cell in excel_row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if excel_row[5].value == "是" or excel_row[6].value == "否":
            for cell in excel_row:
                cell.fill = warn_fill
    sheet.freeze_panes = "A2"

    issue_sheet = workbook.create_sheet("需注意")
    issue_sheet.append(["类型", "scene_id / pair_id", "详情"])
    for cell in issue_sheet[1]:
        cell.fill = head_fill
        cell.font = head_font
    for sid, info in summary["scene_source_conflicts"].items():
        issue_sheet.append(["同场景 source 冲突", sid, json.dumps(info, ensure_ascii=False)[:600]])
    for pid, batch, reason in summary["quality_local_failed"]:
        issue_sheet.append(["自动质量未通过", pid, f"{batch} | {reason}"])
    for pid in summary["deleted_pair_ids"]:
        issue_sheet.append(["已删除（不进入训练）", pid, "人工审计不通过"])
    issue_sheet.column_dimensions["A"].width = 22
    issue_sheet.column_dimensions["B"].width = 46
    issue_sheet.column_dimensions["C"].width = 90
    for excel_row in issue_sheet.iter_rows(min_row=2):
        for cell in excel_row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    return workbook


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.parse_args()
    print(json.dumps(build(), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
