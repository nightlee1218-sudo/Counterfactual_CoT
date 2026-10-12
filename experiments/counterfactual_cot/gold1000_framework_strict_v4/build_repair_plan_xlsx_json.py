#!/usr/bin/env python3
"""Build an audited repair-plan XLSX and isolated JSONL/JSON plan.

No image generation, manifest replacement, pair deletion, or API call happens
here. The plan preserves the original COT/intervention fields and adds proposed
source/target prompt guidance for later human approval.
"""
from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

TAIL = (
    " The edit must be obvious when compared side by side. Keep a native 1024x1024 square, full-bleed scene. "
    "Use the same camera, lighting, background and composition. Do not introduce a frame, border, blank margin, collage, "
    "card, poster, page, mockup, text, logo, watermark or UI."
)


def load_rows(path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["问题汇总"]
    headers = [c.value for c in ws[1]]
    return [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]


def load_manifest(path: Path) -> dict[str, dict[str, Any]]:
    return {r["pair_id"]: r for r in (json.loads(line) for line in path.open(encoding="utf-8") if line.strip())}


def first_entity(pair: dict[str, Any]) -> dict[str, Any]:
    eid = pair["intervention"]["entity_id"]
    return next(e for e in pair["source"]["cot"]["entities"] if e["id"] == eid)


def phrase(entity: dict[str, Any]) -> str:
    attrs = entity.get("attributes", {}) or {}
    values = [str(attrs[k]).replace("_", " ") for k in ("color", "material", "pattern") if attrs.get(k)]
    values.append(str(entity.get("category", "object")).replace("_", " "))
    return " ".join(values)


def contains_any(text: str, values: tuple[str, ...]) -> bool:
    lower = text.lower()
    return any(v.lower() in lower for v in values)


def plan_for(pair: dict[str, Any], issue: str, user_source: bool, scene_pair_count: int) -> dict[str, str]:
    iv = pair.get("intervention", {})
    subtype = str(pair.get("subtype", ""))
    field = str(iv.get("field", ""))
    old = iv.get("old_value")
    new = iv.get("new_value")
    entity = first_entity(pair)
    obj = phrase(entity)
    issue_l = issue.lower()
    source_needed = user_source or contains_any(issue, ("source", "原图", "原source", "source图", "source有", "source没有", "重新生成source", "重新设计source"))
    if contains_any(issue, ("空间不足", "画面拥挤", "其他object", "无关object", "无关物体", "挤", "变短", "位置变化", "位置变动", "影响其他")):
        source_needed = True
    source_scope = "source+target" if source_needed else "target-only"
    source_measure = "保留原始 COT/intervention；重生成时只改善布局可见性，不改变实体类别、数量、属性或关系字段。"
    if source_needed:
        source_measure += " source必须完整1024x1024、无框/无抠图；目标实体数量与old_value一致；只保留任务所需对象；为目标动作预留明确空白；共享scene的其他pair需一并评估。"
    else:
        source_measure = "不修改source；直接基于最终manifest当前source生成target。"

    if field == "count" or subtype in {"1_to_2", "2_to_1", "2_to_3", "3_to_2", "3_to_1", "1_to_3"}:
        if isinstance(old, int) and isinstance(new, int):
            if new > old:
                target = f"在同一现有{obj}区域/同一摞内部增加恰好{new-old}个individual实例；不创建第二摞、第二排、第二层或新区域；原{old}个实例保持原位和细节；最终恰好{new}个。"
            else:
                target = f"从同一现有{obj}区域/同一摞内部移除恰好{old-new}个individual实例；幸存者原位、原尺寸、原细节；最终恰好{new}个；不移动或重排无关物体。"
        else:
            target = f"仅改变{obj}的个体数量从{old}到{new}，以COT中的个体语义计数；不改变位置、尺寸、方向、属性或无关物体。"
    elif subtype in {"absolute_horizontal", "absolute_vertical"}:
        target = f"仅将{obj}从{old}移动到{new}；保持数量、身份、尺寸、方向和细节；清除旧位置残留；所有无关物体、背景和结构完全不动。"
    elif subtype in {"occlusion_state", "containment"}:
        target = f"仅改变{obj}与指定参照物的{field}关系；使用已有物体，不新增或删除任何物体；关系变化必须清晰可见；容器/遮挡物自身保持原位、形状、尺寸和属性，除非它就是intervention.entity_id。"
    elif subtype in {"material", "pattern", "functional_state"} or "attribute" in field:
        target = f"仅改变{obj}的{field}从{old}到{new}；变化必须在目标物体表面/状态上肉眼明显；保持位置、数量、轮廓、身份和其他属性不变。"
    else:
        target = f"仅执行 intervention.field={field} 的单原子变化：{old} -> {new}；其他实体、属性、关系、位置、数量和背景全部保持。"

    if contains_any(issue, ("蜡", "wax", "流蜡")):
        target += "幸存蜡烛必须保留source中的滴蜡、流蜡、污渍和燃烧细节，不得用干净新蜡烛替换。"
    if contains_any(issue, ("不明显", "不明确", "看不出", "肉眼")):
        target += "不得用微小或模糊变化；差异必须在并排比较中一眼可辨。"
    if contains_any(issue, ("悬空", "漂", "切割", "方向", "旋转", "放倒", "挤", "变短")):
        target += "不得悬空、穿透、切割、缩短、旋转或改变任何无关物体。"

    return {
        "repair_scope": source_scope,
        "source_measure": source_measure,
        "target_measure": target + TAIL,
        "preservation_measure": "只允许 intervention.entity_id 的一个字段发生变化；must_preserve 实体和背景必须保持。",
        "validation_measure": "人工复核source old状态、target new状态、个体数量、无新增/消失物体、无关物体位置/尺寸/方向、mask局部性；不合格则不融合。",
        "shared_scene_warning": f"source_scene_id共享{scene_pair_count}个pair；source变更会影响同scene全部pair。" if scene_pair_count > 1 else "该source_scene_id无其他pair共享。",
    }


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--analysis-xlsx", required=True, type=Path)
    ap.add_argument("--manifest", required=True, type=Path)
    ap.add_argument("--output-xlsx", required=True, type=Path)
    ap.add_argument("--output-jsonl", required=True, type=Path)
    ap.add_argument("--output-json", required=True, type=Path)
    args = ap.parse_args()
    rows = load_rows(args.analysis_xlsx)
    manifest = load_manifest(args.manifest)
    scene_counts = Counter(str(r.get("source_scene_id", "")) for r in manifest.values())
    plan_rows = []
    seen = set()
    for row in rows:
        pid = row.get("pair_id_manifest") or row.get("pair_id_raw")
        if not pid or pid in seen or pid not in manifest:
            continue
        seen.add(pid)
        pair = manifest[pid]
        issue = str(row.get("用户审计问题", ""))
        user_source = str(row.get("用户明确标记source重生成", "")) == "是"
        plan = plan_for(pair, issue, user_source, scene_counts[str(pair.get("source_scene_id", ""))])
        plan_rows.append({
            "pair_id": pid, "source_scene_id": pair.get("source_scene_id"), "subtype": pair.get("subtype"),
            "intervention": pair.get("intervention"), "source_cot_text": pair.get("source", {}).get("cot_text"),
            "target_cot_text": pair.get("target", {}).get("cot_text"), "original_target_prompt": pair.get("target", {}).get("prompt"),
            "user_issue": issue, "user_source_marked": user_source, "scene_pair_count": scene_counts[str(pair.get("source_scene_id", ""))],
            **plan, "approval_status": "待人工审核", "generation_status": "未生成", "merge_status": "未融合",
        })

    args.output_jsonl.parent.mkdir(parents=True, exist_ok=True)
    with args.output_jsonl.open("w", encoding="utf-8") as f:
        for row in plan_rows:
            f.write(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n")
    args.output_json.write_text(json.dumps({"schema_version":"counterfactual_repair_plan_v1","source_manifest":str(args.manifest),"pair_count":len(plan_rows),"plans":plan_rows},ensure_ascii=False,indent=2),encoding="utf-8")

    wb = Workbook(); ws = wb.active; ws.title = "修补计划"
    headers = ["序号","pair_id","source_scene_id","scene_pair_count","subtype","intervention_field","old_value","new_value","用户问题","用户标记source","修补范围","Source修改措施","Target修改措施","must_preserve/单原子约束","验收检查","审批状态","生成状态","融合状态"]
    ws.append(headers)
    for i,row in enumerate(plan_rows,1):
        iv=row["intervention"] or {}
        ws.append([i,row["pair_id"],row["source_scene_id"],row["scene_pair_count"],row["subtype"],iv.get("field"),iv.get("old_value"),iv.get("new_value"),row["user_issue"],"是" if row["user_source_marked"] else "否",row["repair_scope"],row["source_measure"],row["target_measure"],row["preservation_measure"],row["validation_measure"],row["approval_status"],row["generation_status"],row["merge_status"]])
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    for c in ws[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1F4E78"); c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(vertical="top",wrap_text=True)
    widths=[8,60,42,15,22,30,15,15,65,15,18,80,100,70,80,16,16,16]
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    table=Table(displayName="RepairPlan",ref=f"A1:R{ws.max_row}"); table.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True,showColumnStripes=False); ws.add_table(table)
    # Summary sheets
    stats=wb.create_sheet("统计"); stats.append(["统计项","数量"]); stats.append(["去重pair数",len(plan_rows)])
    for scope in ("source+target","target-only"):
        stats.append([scope,sum(r["repair_scope"]==scope for r in plan_rows)])
    stats.append([]); stats.append(["subtype","数量"])
    for k,v in sorted(Counter(r["subtype"] for r in plan_rows).items()): stats.append([k,v])
    for c in stats[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1F4E78")
    stats.column_dimensions['A'].width=35; stats.column_dimensions['B'].width=15
    wb.save(args.output_xlsx)
    print(json.dumps({"xlsx":str(args.output_xlsx),"jsonl":str(args.output_jsonl),"json":str(args.output_json),"pair_count":len(plan_rows),"source_target":sum(r['repair_scope']=='source+target' for r in plan_rows),"target_only":sum(r['repair_scope']=='target-only' for r in plan_rows)},ensure_ascii=False,indent=2))

if __name__=='__main__': main()
