#!/usr/bin/env python3
"""Consolidate per-pair repair plans into source-scene and target plans.

This planning-only command does not call APIs, generate images, modify the
1000-pair manifest, or change COT. It makes shared-source spillover explicit and
flags group-level entity selectors that need a human semantic decision.
"""
from __future__ import annotations

import argparse
import json
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


def load_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.open(encoding="utf-8") if line.strip()]


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--detailed-jsonl", required=True, type=Path)
    ap.add_argument("--manifest", required=True, type=Path)
    ap.add_argument("--output-xlsx", required=True, type=Path)
    ap.add_argument("--output-json", required=True, type=Path)
    args = ap.parse_args()

    plans = {r["pair_id"]: r for r in load_jsonl(args.detailed_jsonl)}
    manifest = {r["pair_id"]: r for r in load_jsonl(args.manifest)}
    by_scene: dict[str, list[str]] = defaultdict(list)
    for pid, record in manifest.items():
        by_scene[str(record["source_scene_id"])].append(pid)

    scene_rows: list[dict[str, Any]] = []
    target_rows: list[dict[str, Any]] = []
    selector_rows: list[dict[str, Any]] = []
    for scene_id in sorted({str(r["source_scene_id"]) for r in plans.values()}):
        all_members = sorted(by_scene.get(scene_id, []))
        reported = sorted(pid for pid, r in plans.items() if str(r["source_scene_id"]) == scene_id)
        source_plans = [plans[pid] for pid in reported if plans[pid].get("repair_scope") == "source+target"]
        source_required = bool(source_plans)
        measures = list(dict.fromkeys(str(r.get("source_measure", "")).strip() for r in source_plans if str(r.get("source_measure", "")).strip()))
        base_prompts = list(dict.fromkeys(str(r.get("original_source_prompt", "")) for r in source_plans if r.get("original_source_prompt")))
        scene_rows.append({
            "source_scene_id": scene_id,
            "all_scene_pair_ids": all_members,
            "reported_problem_pair_ids": reported,
            "source_required": source_required,
            "source_plan_count": len(source_plans),
            "source_prompt_base_count": len(base_prompts),
            "source_constraint_count": len(measures),
            "source_prompt_conflict_review": "YES" if len(base_prompts) > 1 or len(measures) > 1 else "NO",
            "source_constraints_to_merge": measures,
            "source_prompt_status": "待人工合并约束后生成" if source_required else "不重生成source",
            "spillover_pair_ids": sorted(set(all_members) - set(reported)) if source_required else [],
        })
        for pid in all_members if source_required else reported:
            original = manifest[pid]
            plan = plans.get(pid)
            spillover = plan is None
            if spillover:
                scope = "source+target spillover"
                target_measure = "Source scene will change; perform the original target intervention exactly once using the new source. Preserve all unrelated entities and counts."
                entity_count = None
                ambiguity = "需按原COT复核"
            else:
                scope = str(plan.get("repair_scope"))
                target_measure = str(plan.get("target_measure", ""))
                entity_count = plan.get("entity_group_count")
                ambiguity = str(plan.get("entity_count_ambiguity", "低"))
            iv = original.get("intervention", {})
            target_rows.append({
                "pair_id": pid, "source_scene_id": scene_id, "repair_scope": scope,
                "repair_role": "source_scene_spillover" if spillover else "reported_problem",
                "subtype": original.get("subtype"), "entity_id": iv.get("entity_id"),
                "entity_group_count": entity_count, "entity_count_ambiguity": ambiguity,
                "intervention_field": iv.get("field"), "old_value": iv.get("old_value"), "new_value": iv.get("new_value"),
                "target_cot_text": original.get("target", {}).get("cot_text"),
                "original_target_prompt": original.get("target", {}).get("prompt"),
                "target_measure": target_measure, "generation_status": "未生成", "approval_status": "待人工审核",
            })
            if plan and plan.get("entity_count_ambiguity") == "高":
                selector_rows.append({
                    "pair_id": pid, "source_scene_id": scene_id, "entity_id": iv.get("entity_id"),
                    "group_count": entity_count, "subtype": original.get("subtype"), "field": iv.get("field"),
                    "old_value": iv.get("old_value"), "new_value": iv.get("new_value"),
                    "decision_needed": "先确定是整个entity group操作，还是单个实例操作；不能把entity_id后缀当作视觉索引。",
                    "safe_default": "group-level operation；若用户明确只改一个实例，必须补充leftmost/frontmost/nearest等视觉selector后再生成。",
                    "status": "待人工确认",
                })

    payload = {
        "schema_version": "counterfactual_scene_repair_plan_v1",
        "detailed_plan": str(args.detailed_jsonl), "manifest": str(args.manifest),
        "scene_count": len(scene_rows), "reported_target_count": len(plans),
        "expanded_target_count": len(target_rows), "source_scene_rows": scene_rows,
        "target_rows": target_rows, "entity_selector_rows": selector_rows,
    }
    args.output_json.parent.mkdir(parents=True, exist_ok=True)
    args.output_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    wb = Workbook(); scenes = wb.active; scenes.title = "SourceScene计划"
    scene_headers = ["source_scene_id", "all_scene_pair_ids", "reported_problem_pair_ids", "source_required", "source_plan_count", "source_prompt_base_count", "source_constraint_count", "source_prompt_conflict_review", "source_constraints_to_merge", "source_prompt_status", "spillover_pair_ids"]
    scenes.append(scene_headers)
    for r in scene_rows: scenes.append([r.get(h, "") if not isinstance(r.get(h), list) else ";".join(r[h]) for h in scene_headers])
    targets = wb.create_sheet("Target计划")
    target_headers = ["pair_id", "source_scene_id", "repair_scope", "repair_role", "subtype", "entity_id", "entity_group_count", "entity_count_ambiguity", "intervention_field", "old_value", "new_value", "target_cot_text", "original_target_prompt", "target_measure", "approval_status", "generation_status"]
    targets.append(target_headers)
    for r in target_rows: targets.append([r.get(h, "") for h in target_headers])
    selectors = wb.create_sheet("EntitySelector待确认")
    selector_headers = ["pair_id", "source_scene_id", "entity_id", "group_count", "subtype", "field", "old_value", "new_value", "decision_needed", "safe_default", "status"]
    selectors.append(selector_headers)
    for r in selector_rows: selectors.append([r.get(h, "") for h in selector_headers])
    stats = wb.create_sheet("统计"); stats.append(["统计项", "数量"]); stats.append(["source_scene数量", len(scene_rows)]); stats.append(["source需要重生成scene", sum(r["source_required"] for r in scene_rows)]); stats.append(["报告问题target数", len(plans)]); stats.append(["扩展后target数（含spillover）", len(target_rows)]); stats.append(["Entity selector待确认", len(selector_rows)])
    for ws in (scenes, targets, selectors, stats):
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E78"); c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for c in row: c.alignment = Alignment(vertical="top", wrap_text=True)
    for ws in (scenes, targets, selectors):
        for i in range(1, ws.max_column + 1): ws.column_dimensions[get_column_letter(i)].width = 32
    for ws, name in ((scenes, "SourceScenePlan"), (targets, "TargetPlan"), (selectors, "EntitySelectorPlan")):
        if ws.max_row > 1:
            table = Table(displayName=name, ref=f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"); table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False); ws.add_table(table)
    args.output_xlsx.parent.mkdir(parents=True, exist_ok=True); wb.save(args.output_xlsx)
    print(json.dumps({"output_xlsx": str(args.output_xlsx), "output_json": str(args.output_json), "source_scenes": len(scene_rows), "source_required_scenes": sum(r["source_required"] for r in scene_rows), "reported_targets": len(plans), "expanded_targets": len(target_rows), "entity_selector_review": len(selector_rows)}, ensure_ascii=False, indent=2))


if __name__ == "__main__": main()
